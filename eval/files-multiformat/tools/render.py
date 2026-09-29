"""Renderers: one structured, fully synthetic document -> one file of a given type.

A document ("doc") is a dict:
  title, lang ("zh"/"en"), blocks: [("h", text) | ("p", text) | ("kv", [(k, v), ...]) | ("ul", [text, ...])
  | ("table", {"columns": [...], "rows": [[...]], "caption": str})], optional slides: [{"title", "bullets", "image": bool}],
  optional sheets (spreadsheets), optional record (json/xml), optional visual (image style hints).
Every renderer draws a small "合成数据" (synthetic data) mark where a reader can see it.

Only non-ML code: Pillow, python-docx, XlsxWriter, python-pptx, odfpy, PyMuPDF, zipfile/plistlib/email,
and the Mac's textutil/sips/plutil; video frames are encoded with the ffmpeg binary from imageio-ffmpeg.
"""

from __future__ import annotations

import base64
import datetime as dt
import html
import io
import math
import os
import plistlib
import random
import re
import shutil
import subprocess
import tempfile
import zipfile

from PIL import Image, ImageDraw, ImageFilter, ImageFont

MARK = "合成数据"
MARK_EN = "SYNTHETIC DATA"
FIXED_TIME = dt.datetime(2026, 10, 1, 9, 0, 0)

CJK_TTC = "/System/Library/Fonts/Hiragino Sans GB.ttc"  # face 0 = W3, face 1 = W6
CJK_TTF = "/Library/Fonts/Arial Unicode.ttf"  # TrueType outlines: embeddable (subset) in PDFs
MONO = "/System/Library/Fonts/Menlo.ttc"

_font_cache: dict = {}
OVERFLOWS: list[str] = []  # text that did not fit where it was drawn (the builder refuses such files)


def font(size: int, bold: bool = False, mono: bool = False) -> ImageFont.FreeTypeFont:
    key = (size, bold, mono)
    if key not in _font_cache:
        if mono:
            _font_cache[key] = ImageFont.truetype(MONO, size, index=0)
        else:
            _font_cache[key] = ImageFont.truetype(CJK_TTC, size, index=1 if bold else 0)
    return _font_cache[key]


# ----------------------------------------------------------------------------- text helpers

def mark_for(doc: dict) -> str:
    return MARK if doc.get("lang", "zh") == "zh" else f"{MARK_EN} · {MARK}"


def cell_str(v) -> str:
    if isinstance(v, float):
        return f"{v:,.2f}" if abs(v) >= 1000 else (f"{v:g}")
    if isinstance(v, int) and not isinstance(v, bool):
        return f"{v:,}" if abs(v) >= 10000 else str(v)
    return str(v)


def blocks_to_lines(doc: dict, kv_sep: str = "：") -> list[str]:
    """Plain-text reading of a document (also the canonical ground-truth text)."""
    sep = kv_sep if doc.get("lang", "zh") == "zh" else ": "
    out = [doc["title"], ""]
    for b in doc["blocks"]:
        kind = b[0]
        if kind == "h":
            out += [b[1]]
        elif kind == "p":
            out += [b[1]]
        elif kind == "kv":
            out += [f"{k}{sep}{v}" for k, v in b[1]]
        elif kind == "ul":
            out += [f"- {x}" for x in b[1]]
        elif kind == "table":
            t = b[1]
            if t.get("caption"):
                out.append(t["caption"])
            out.append("\t".join(t["columns"]))
            out += ["\t".join(cell_str(c) for c in r) for r in t["rows"]]
        out.append("")
    return out


def plain_text(doc: dict) -> str:
    return "\n".join(blocks_to_lines(doc)).strip() + "\n"


def wrap(text: str, fnt: ImageFont.FreeTypeFont, width: int) -> list[str]:
    tokens = re.findall(r"[A-Za-z0-9_\-./:@%#+()'\",$]+|\s+|.", text)
    lines, cur = [], ""
    for tok in tokens:
        cand = cur + tok
        if fnt.getlength(cand) <= width or not cur:
            cur = cand
        else:
            lines.append(cur.rstrip())
            cur = tok.lstrip()
    if cur.strip():
        lines.append(cur.rstrip())
    return lines or [""]


# ----------------------------------------------------------------------------- raster pages / cards

PAPER = (253, 252, 248)
INK = (30, 32, 36)
MUTED = (120, 124, 132)
RED = (200, 60, 50)


def draw_mark(d: ImageDraw.ImageDraw, w: int, h: int, text: str = MARK, size: int = 18, color=RED, corner="br"):
    f = font(size)
    tw = f.getlength(text)
    x = w - tw - 16 if corner.endswith("r") else 16
    y = h - size - 14 if corner.startswith("b") else 12
    d.text((x, y), text, font=f, fill=color)


def render_page_images(doc: dict, width: int = 1240, height: int = 1754, margin: int = 110, scale: float = 1.0) -> list[Image.Image]:
    """A4-like pages (about 150 dpi) drawn from the blocks; used for scans, photos and TIFFs."""
    s = scale
    W, H, M = int(width * s), int(height * s), int(margin * s)
    fz = {"title": int(40 * s), "h": int(30 * s), "p": int(24 * s), "small": int(18 * s)}
    pages: list[Image.Image] = []

    def new_page():
        im = Image.new("RGB", (W, H), PAPER)
        return im, ImageDraw.Draw(im), M

    im, d, y = new_page()

    def need(hh):
        nonlocal im, d, y
        if y + hh > H - M:
            draw_mark(d, W, H, mark_for(doc), size=fz["small"])
            pages.append(im)
            im, d, y = new_page()

    if doc.get("letterhead"):
        f = font(fz["small"])
        d.text((M, y - int(50 * s)), doc["letterhead"], font=f, fill=MUTED)
        d.line((M, y - int(20 * s), W - M, y - int(20 * s)), fill=(180, 180, 180), width=2)
    for line in wrap(doc["title"], font(fz["title"], True), W - 2 * M):
        need(fz["title"] + 20)
        tw = font(fz["title"], True).getlength(line)
        d.text(((W - tw) / 2, y), line, font=font(fz["title"], True), fill=INK)
        y += fz["title"] + 18
    y += 20
    for b in doc["blocks"]:
        kind = b[0]
        if kind == "pagebreak":  # only the paged picture renderers honour it
            need(H)
            continue
        if kind == "h":
            need(fz["h"] + 30)
            y += 10
            d.text((M, y), b[1], font=font(fz["h"], True), fill=INK)
            y += fz["h"] + 16
        elif kind in ("p", "ul", "kv"):
            items = [b[1]] if kind == "p" else ([f"• {x}" for x in b[1]] if kind == "ul" else
                                                [f"{k}{'：' if doc.get('lang','zh')=='zh' else ': '}{v}" for k, v in b[1]])
            for it in items:
                for line in wrap(it, font(fz["p"]), W - 2 * M):
                    need(fz["p"] + 12)
                    d.text((M, y), line, font=font(fz["p"]), fill=INK)
                    y += fz["p"] + 12
                y += 6
            y += 8
        elif kind == "table":
            t = b[1]
            cols = t["columns"]
            rows = [[cell_str(c) for c in r] for r in t["rows"]]
            tsz = fz["small"] + 2
            avail = W - 2 * M
            while True:
                fnt = font(tsz)
                widths = [max(font(tsz, True).getlength(str(c)) for c in [cols[i]] + [r[i] for r in rows]) + 24 for i in range(len(cols))]
                if sum(widths) <= avail or tsz <= 12:
                    break
                tsz -= 1
            if sum(widths) > avail:
                OVERFLOWS.append(f"page table too wide: {cols}")
            rh = tsz + 22
            if t.get("caption"):
                need(rh)
                d.text((M, y), t["caption"], font=font(fz["small"] + 2, True), fill=INK)
                y += rh
            for ri, r in enumerate([cols] + rows):
                need(rh)
                x = M
                for ci, c in enumerate(r):
                    d.rectangle((x, y, x + widths[ci], y + rh), outline=(90, 90, 90), width=1,
                                fill=(236, 238, 242) if ri == 0 else None)
                    d.text((x + 8, y + 9), str(c), font=font(tsz, ri == 0), fill=INK)
                    x += widths[ci]
                y += rh
            y += 20
    if doc.get("signature"):
        need(160)
        y += 30
        sig_w = font(fz["p"]).getlength(doc["signature"])
        sx = W - M - sig_w - (230 * s if doc.get("stamp") else 0)
        d.text((sx, y), doc["signature"], font=font(fz["p"]), fill=INK)
        if doc.get("stamp"):
            cx, cy, r = int(W - M - 110 * s), int(y + 10 * s), int(95 * s)
            d.ellipse((cx - r, cy - r, cx + r, cy + r), outline=(210, 40, 40), width=int(5 * s))
            sf = font(int(22 * s), True)
            tw = sf.getlength(doc["stamp"])
            d.text((cx - tw / 2, cy - 12 * s), doc["stamp"], font=sf, fill=(210, 40, 40))
            d.text((cx - 12 * s, cy - 60 * s), "★", font=font(int(26 * s)), fill=(210, 40, 40))
    draw_mark(d, W, H, mark_for(doc), size=fz["small"])
    pages.append(im)
    return pages


def scan_effect(im: Image.Image, rng: random.Random, angle: float | None = None, noise: float = 18.0) -> Image.Image:
    """Photocopier/scanner look: grey paper, skew, speckle noise, slight blur."""
    im = im.convert("L")
    ang = angle if angle is not None else rng.uniform(-2.2, 2.2)
    im = im.rotate(ang, resample=Image.BICUBIC, expand=False, fillcolor=235)
    noise_img = Image.effect_noise(im.size, noise).convert("L")
    im = Image.blend(im, noise_img, 0.12)
    im = im.point(lambda v: min(255, int(v * 0.93 + 8)))
    im = im.filter(ImageFilter.GaussianBlur(0.6))
    d = ImageDraw.Draw(im)
    for _ in range(120):
        x, y = rng.randrange(im.width), rng.randrange(im.height)
        d.point((x, y), fill=rng.randrange(40, 120))
    return im


def photo_effect(im: Image.Image, rng: random.Random, bg=(92, 78, 64)) -> Image.Image:
    """A phone photo of a paper/label lying on a desk: margin, rotation, shadow, warm light, noise."""
    w, h = im.size
    canvas = Image.new("RGB", (int(w * 1.18), int(h * 1.18)), bg)
    shadow = Image.new("RGBA", (w, h), (0, 0, 0, 110))
    ang = rng.uniform(-4, 4)
    sh = shadow.rotate(ang, expand=True)
    pic = im.convert("RGBA").rotate(ang, expand=True, resample=Image.BICUBIC)
    ox, oy = (canvas.width - pic.width) // 2, (canvas.height - pic.height) // 2
    canvas.paste(sh, (ox + 14, oy + 18), sh)
    canvas = canvas.filter(ImageFilter.GaussianBlur(2))
    canvas.paste(pic, (ox, oy), pic)
    # uneven light
    grad = Image.linear_gradient("L").resize(canvas.size).rotate(rng.choice([0, 90, 180, 270]))
    warm = Image.new("RGB", canvas.size, (255, 236, 205))
    canvas = Image.composite(canvas, Image.blend(canvas, warm, 0.25), grad)
    noise_img = Image.effect_noise(canvas.size, 10).convert("RGB")
    canvas = Image.blend(canvas, noise_img, 0.05)
    return canvas


def render_slide_image(title: str, bullets: list[str], w: int = 1280, h: int = 720, theme: int = 0,
                       footer: str = MARK, table: dict | None = None, big: str | None = None) -> Image.Image:
    themes = [((255, 255, 255), (32, 64, 120), (40, 44, 52)), ((20, 30, 48), (120, 200, 255), (235, 238, 245)),
              ((248, 244, 236), (160, 70, 40), (50, 40, 32)), ((236, 246, 240), (30, 110, 80), (30, 40, 36))]
    bg, accent, fg = themes[theme % len(themes)]
    im = Image.new("RGB", (w, h), bg)
    d = ImageDraw.Draw(im)
    d.rectangle((0, 0, w, 10), fill=accent)
    y = 50
    for line in wrap(title, font(46, True), w - 140):
        d.text((70, y), line, font=font(46, True), fill=accent)
        y += 60
    y += 20
    if big:
        f = font(96, True)
        tw = f.getlength(big)
        d.text(((w - tw) / 2, y + 20), big, font=f, fill=accent)
        y += 150
    for b in bullets:
        lines = wrap(b, font(30), w - 200)
        d.ellipse((78, y + 14, 90, y + 26), fill=accent)
        for i, line in enumerate(lines):
            d.text((110, y), line, font=font(30), fill=fg)
            y += 44
        y += 10
    if table:
        cols, rows = table["columns"], [[cell_str(c) for c in r] for r in table["rows"]]
        fnt = font(26)
        cw = (w - 160) / len(cols)
        rh = 48
        for ri, r in enumerate([cols] + rows):
            for ci, c in enumerate(r):
                x0 = 80 + ci * cw
                d.rectangle((x0, y, x0 + cw, y + rh), outline=accent, width=2,
                            fill=(accent if ri == 0 else None))
                if font(26, ri == 0).getlength(str(c)) > cw - 16:
                    OVERFLOWS.append(f"slide cell too wide: {c}")
                d.text((x0 + 12, y + 10), str(c), font=font(26, ri == 0), fill=(bg if ri == 0 else fg))
            y += rh
    f = font(18)
    d.text((w - f.getlength(footer) - 24, h - 36), footer, font=f, fill=RED)
    return im


def render_card(title: str, lines: list[str], w: int = 1000, bg=(255, 255, 255), accent=(40, 90, 160), fg=INK,
                fields: list[tuple[str, str]] | None = None, table: dict | None = None, title_size: int = 40,
                line_size: int = 28, footer: str = MARK, pad: int = 60, center_title: bool = False,
                min_h: int = 0) -> Image.Image:
    """Poster / notice / label / receipt style card. Height grows with content."""
    items = []
    tf = font(title_size, True)
    for line in wrap(title, tf, w - 2 * pad):
        items.append(("t", line))
    items.append(("gap", 16))
    for ln in lines:
        for sub in wrap(ln, font(line_size), w - 2 * pad):
            items.append(("l", sub))
        items.append(("gap", 6))
    if fields:
        items.append(("gap", 10))
        for k, v in fields:
            items.append(("f", (k, v)))
    if table:
        items.append(("gap", 16))
        items.append(("table", table))
    # measure
    hh = pad
    tbl_rh = line_size + 22
    for kind, val in items:
        if kind == "t":
            hh += title_size + 16
        elif kind == "l" or kind == "f":
            hh += line_size + 14
        elif kind == "gap":
            hh += val
        elif kind == "table":
            hh += tbl_rh * (len(val["rows"]) + 1) + 10
    hh += pad + 30
    H = max(hh, min_h)
    im = Image.new("RGB", (w, H), bg)
    d = ImageDraw.Draw(im)
    d.rectangle((0, 0, w, 12), fill=accent)
    y = pad
    for kind, val in items:
        if kind == "t":
            x = (w - tf.getlength(val)) / 2 if center_title else pad
            d.text((x, y), val, font=tf, fill=accent)
            y += title_size + 16
        elif kind == "l":
            d.text((pad, y), val, font=font(line_size), fill=fg)
            y += line_size + 14
        elif kind == "f":
            k, v = val
            d.text((pad, y), f"{k}", font=font(line_size, True), fill=MUTED if bg == (255, 255, 255) else fg)
            vx = pad + max(220, font(line_size, True).getlength(k) + 30)
            if vx + font(line_size).getlength(str(v)) > w - 10:
                OVERFLOWS.append(f"card field too wide: {k}={v}")
            d.text((vx, y), str(v), font=font(line_size), fill=fg)
            y += line_size + 14
        elif kind == "gap":
            y += val
        elif kind == "table":
            cols, rows = val["columns"], [[cell_str(c) for c in r] for r in val["rows"]]
            tsz = line_size - 4
            while True:
                widths = [max(font(tsz, True).getlength(str(c)) for c in [cols[i]] + [r[i] for r in rows]) + 28 for i in range(len(cols))]
                if sum(widths) <= w - 2 * pad or tsz <= 12:
                    break
                tsz -= 1
            if sum(widths) > w - 2 * pad:
                OVERFLOWS.append(f"card table too wide: {cols}")
            for ri, r in enumerate([cols] + rows):
                x = pad
                for ci, c in enumerate(r):
                    d.rectangle((x, y, x + widths[ci], y + tbl_rh), outline=(150, 150, 150),
                                fill=(accent if ri == 0 else (bg if ri % 2 else (246, 247, 250))))
                    d.text((x + 10, y + 10), str(c), font=font(tsz, ri == 0),
                           fill=((255, 255, 255) if ri == 0 else fg))
                    x += widths[ci]
                y += tbl_rh
            y += 10
    f = font(18)
    d.text((w - f.getlength(footer) - 18, H - 34), footer, font=f, fill=RED)
    return im


def render_chart(title: str, series: dict, xlabels: list[str], kind: str = "line", w: int = 1200, h: int = 760,
                 ylabel: str = "", notes: list[str] | None = None, value_labels: bool = True) -> Image.Image:
    """Simple line/bar chart with axis ticks, legend and optional annotations (all text is drawn)."""
    im = Image.new("RGB", (w, h), (255, 255, 255))
    d = ImageDraw.Draw(im)
    d.text((60, 28), title, font=font(34, True), fill=INK)
    L, T, R, B = 110, 110, w - 60, h - 140
    allv = [v for s in series.values() for v in s]
    vmin = 0 if kind == "bar" else min(allv)
    vmax = max(allv)
    span = (vmax - vmin) or 1
    vmin_p, vmax_p = vmin - (0 if kind == "bar" else span * 0.1), vmax + span * 0.15
    def Y(v):
        return B - (v - vmin_p) / (vmax_p - vmin_p) * (B - T)
    d.line((L, T, L, B), fill=INK, width=2)
    d.line((L, B, R, B), fill=INK, width=2)
    for i in range(5):
        v = vmin_p + (vmax_p - vmin_p) * i / 4
        yy = Y(v)
        d.line((L, yy, R, yy), fill=(228, 228, 232), width=1)
        lab = f"{v:.2f}" if span < 10 else f"{v:,.0f}"
        d.text((L - 10 - font(18).getlength(lab), yy - 10), lab, font=font(18), fill=MUTED)
    if ylabel:
        d.text((20, T - 40), ylabel, font=font(20), fill=MUTED)
    colors = [(40, 90, 170), (220, 110, 40), (40, 150, 90), (160, 60, 150)]
    n = len(xlabels)
    step = (R - L) / max(n, 1)
    for i, lab in enumerate(xlabels):
        xx = L + step * (i + 0.5)
        d.text((xx - font(18).getlength(lab) / 2, B + 10), lab, font=font(18), fill=INK)
    k = len(series)
    for si, (name, vals) in enumerate(series.items()):
        c = colors[si % len(colors)]
        if kind == "line":
            pts = [(L + step * (i + 0.5), Y(v)) for i, v in enumerate(vals)]
            d.line(pts, fill=c, width=4)
            for (px, py), v in zip(pts, vals):
                d.ellipse((px - 5, py - 5, px + 5, py + 5), fill=c)
                if value_labels:
                    lab = f"{v:g}"
                    d.text((px - font(16).getlength(lab) / 2, py - 28 if si % 2 == 0 else py + 10), lab, font=font(16), fill=c)
        else:
            bw = step * 0.7 / k
            for i, v in enumerate(vals):
                x0 = L + step * i + step * 0.15 + si * bw
                d.rectangle((x0, Y(v), x0 + bw - 4, B), fill=c)
                if value_labels:
                    lab = f"{v:g}" if not isinstance(v, int) else f"{v:,}"
                    d.text((x0 + (bw - 4) / 2 - font(16).getlength(lab) / 2, Y(v) - 24), lab, font=font(16), fill=INK)
    lx = L
    for si, name in enumerate(series):
        c = colors[si % len(colors)]
        d.rectangle((lx, h - 90, lx + 22, h - 70), fill=c)
        d.text((lx + 30, h - 94), name, font=font(20), fill=INK)
        lx += 60 + font(20).getlength(name)
    if notes:
        ny = T + 10
        for nt in notes:
            tw = font(20).getlength(nt)
            d.rectangle((R - tw - 30, ny - 6, R - 6, ny + 28), fill=(255, 248, 225), outline=(220, 180, 80))
            d.text((R - tw - 18, ny), nt, font=font(20), fill=INK)
            ny += 44
    f = font(18)
    d.text((w - f.getlength(MARK) - 16, h - 32), MARK, font=f, fill=RED)
    return im


def render_window(app_title: str, lines: list[str], fields: list[tuple[str, str]] | None = None,
                  table: dict | None = None, w: int = 1280, badge: str | None = None) -> Image.Image:
    """Generic desktop-app/web screenshot: title bar, content card (no real product branding)."""
    body = render_card(lines[0] if lines else app_title, lines[1:], w=w - 80, fields=fields, table=table,
                       title_size=34, line_size=26, footer=MARK)
    H = body.height + 120
    im = Image.new("RGB", (w, H), (236, 238, 242))
    d = ImageDraw.Draw(im)
    d.rectangle((0, 0, w, 52), fill=(250, 250, 252))
    for i, c in enumerate([(236, 95, 90), (245, 190, 80), (98, 197, 84)]):
        d.ellipse((20 + i * 28, 18, 36 + i * 28, 34), fill=c)
    d.text((w / 2 - font(22).getlength(app_title) / 2, 13), app_title, font=font(22), fill=INK)
    im.paste(body, (40, 80))
    if badge:
        bf = font(24, True)
        tw = bf.getlength(badge)
        d.rounded_rectangle((w - tw - 90, 90, w - 50, 134), radius=18, fill=(40, 160, 90))
        d.text((w - tw - 70, 98), badge, font=bf, fill=(255, 255, 255))
    return im


# ----------------------------------------------------------------------------- office-like formats

def write_txt(path: str, text: str, encoding: str = "utf-8", newline: str = "\n"):
    data = text.replace("\n", newline)
    with open(path, "wb") as fh:
        fh.write(data.encode(encoding))


def to_markdown(doc: dict) -> str:
    out = []
    if doc.get("front_matter"):
        out += ["---"] + [f"{k}: {v}" for k, v in doc["front_matter"].items()] + ["---", ""]
    out += [f"# {doc['title']}", ""]
    for b in doc["blocks"]:
        k = b[0]
        if k == "h":
            out += [f"## {b[1]}", ""]
        elif k == "p":
            out += [b[1], ""]
        elif k == "kv":
            out += [f"- **{a}**：{v}" if doc.get("lang", "zh") == "zh" else f"- **{a}**: {v}" for a, v in b[1]] + [""]
        elif k == "ul":
            out += [f"- [ ] {x[4:]}" if x.startswith("[ ] ") else (f"- [x] {x[4:]}" if x.startswith("[x] ") else f"- {x}")
                    for x in b[1]] + [""]
        elif k == "table":
            t = b[1]
            if t.get("caption"):
                out += [f"**{t['caption']}**", ""]
            out += ["| " + " | ".join(t["columns"]) + " |", "|" + "---|" * len(t["columns"])]
            out += ["| " + " | ".join(cell_str(c) for c in r) + " |" for r in t["rows"]] + [""]
        elif k == "code":
            out += ["```" + b[1].get("lang", ""), b[1]["text"], "```", ""]
    out += [f"> {mark_for(doc)}", ""]
    return "\n".join(out)


def to_html(doc: dict, full: bool = True, extra_head: str = "", nav: bool = True) -> str:
    e = html.escape
    body = []
    if nav:
        body.append(f'<header class="top"><span class="brand">{e(doc.get("site", "内部页面"))}</span>'
                    f'<span class="mark">{MARK}</span></header>')
    body.append(f"<main><h1>{e(doc['title'])}</h1>")
    for b in doc["blocks"]:
        k = b[0]
        if k == "h":
            body.append(f"<h2>{e(b[1])}</h2>")
        elif k == "p":
            body.append(f"<p>{e(b[1])}</p>")
        elif k == "kv":
            body.append("<dl>" + "".join(f"<dt>{e(a)}</dt><dd>{e(str(v))}</dd>" for a, v in b[1]) + "</dl>")
        elif k == "ul":
            body.append("<ul>" + "".join(f"<li>{e(x)}</li>" for x in b[1]) + "</ul>")
        elif k == "table":
            t = b[1]
            cap = f"<caption>{e(t['caption'])}</caption>" if t.get("caption") else ""
            head = "".join(f"<th>{e(c)}</th>" for c in t["columns"])
            rows = "".join("<tr>" + "".join(f"<td>{e(cell_str(c))}</td>" for c in r) + "</tr>" for r in t["rows"])
            body.append(f"<table>{cap}<thead><tr>{head}</tr></thead><tbody>{rows}</tbody></table>")
        elif k == "img":
            body.append(f'<figure><img src="{e(b[1]["src"])}" alt="{e(b[1].get("alt", ""))}" width="480">'
                        f'<figcaption>{e(b[1].get("caption", ""))}</figcaption></figure>')
    body.append(f'</main><footer><small>{e(mark_for(doc))} · {e(doc.get("footer", "仅供评测使用"))}</small></footer>')
    if not full:
        return "\n".join(body)
    lang = "zh-CN" if doc.get("lang", "zh") == "zh" else "en"
    css = ("body{font-family:-apple-system,'PingFang SC','Hiragino Sans GB',sans-serif;margin:0;color:#222;background:#fafafa}"
           "header.top{display:flex;justify-content:space-between;padding:10px 24px;background:#1f3b63;color:#fff}"
           ".mark{color:#ffb4a8;font-size:12px}main{max-width:860px;margin:24px auto;background:#fff;padding:24px 32px}"
           "table{border-collapse:collapse}td,th{border:1px solid #bbb;padding:4px 10px}dt{font-weight:600}"
           "footer{text-align:center;color:#888;padding:16px}")
    return (f'<!DOCTYPE html>\n<html lang="{lang}">\n<head>\n<meta charset="utf-8">\n<title>{e(doc["title"])}</title>\n'
            f'<meta name="generator" content="synthetic-eval">\n<style>{css}</style>\n{extra_head}</head>\n<body>\n'
            + "\n".join(body) + "\n</body>\n</html>\n")


def _docx_set_cjk(run, name="PingFang SC"):
    from docx.oxml.ns import qn
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)


def write_docx(path: str, doc: dict):
    import docx
    from docx.shared import Pt
    from docx.oxml.ns import qn
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    cp = d.core_properties
    cp.author = doc.get("author", "synthetic")
    cp.title = doc["title"]
    cp.created = FIXED_TIME
    cp.modified = FIXED_TIME
    cp.last_modified_by = "synthetic"
    cp.comments = MARK
    sec = d.sections[0]
    sec.header.paragraphs[0].text = doc.get("letterhead", "")
    sec.footer.paragraphs[0].text = f"{mark_for(doc)} · {doc.get('footer', '仅供评测使用')}"
    d.add_heading(doc["title"], level=0)
    for b in doc["blocks"]:
        k = b[0]
        if k == "h":
            d.add_heading(b[1], level=1)
        elif k == "p":
            d.add_paragraph(b[1])
        elif k == "kv":
            for a, v in b[1]:
                p = d.add_paragraph()
                r = p.add_run(f"{a}{'：' if doc.get('lang','zh')=='zh' else ': '}")
                r.bold = True
                p.add_run(str(v))
        elif k == "ul":
            for x in b[1]:
                d.add_paragraph(x, style="List Bullet")
        elif k == "table":
            t = b[1]
            if t.get("caption"):
                d.add_paragraph(t["caption"]).runs[0].bold = True
            tb = d.add_table(rows=1, cols=len(t["columns"]))
            tb.style = "Table Grid"
            for i, c in enumerate(t["columns"]):
                tb.rows[0].cells[i].text = c
            for r in t["rows"]:
                cells = tb.add_row().cells
                for i, c in enumerate(r):
                    cells[i].text = cell_str(c)
    if doc.get("signature"):
        d.add_paragraph("")
        d.add_paragraph(doc["signature"])
    mp = d.add_paragraph()
    mr = mp.add_run(mark_for(doc))
    mr.font.size = Pt(8)
    d.save(path)


def textutil_convert(src: str, dst: str, fmt: str):
    subprocess.run(["/usr/bin/textutil", "-convert", fmt, src, "-output", dst], check=True,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)


def write_rtf(path: str, doc: dict):
    with tempfile.TemporaryDirectory() as td:
        src = os.path.join(td, "in.html")
        with open(src, "w", encoding="utf-8") as fh:
            fh.write(to_html(doc, nav=False))
        textutil_convert(src, path, "rtf")


def write_doc(path: str, doc: dict):
    with tempfile.TemporaryDirectory() as td:
        src = os.path.join(td, "in.docx")
        write_docx(src, doc)
        textutil_convert(src, path, "doc")


def write_odt(path: str, doc: dict):
    from odf.opendocument import OpenDocumentText
    from odf.style import Style, TextProperties, TableColumnProperties, ParagraphProperties
    from odf.text import H, P, List, ListItem, Span
    from odf.table import Table, TableColumn, TableRow, TableCell
    from odf import meta, dc
    o = OpenDocumentText()
    o.meta.addElement(dc.Title(text=doc["title"]))
    o.meta.addElement(meta.CreationDate(text=FIXED_TIME.isoformat()))
    bold = Style(name="B", family="text")
    bold.addElement(TextProperties(fontweight="bold", fontweightasian="bold"))
    o.automaticstyles.addElement(bold)
    h1 = Style(name="Heading 1", family="paragraph")
    h1.addElement(TextProperties(fontsize="18pt", fontweight="bold", fontsizeasian="18pt", fontweightasian="bold"))
    o.styles.addElement(h1)
    h2 = Style(name="Heading 2", family="paragraph")
    h2.addElement(TextProperties(fontsize="14pt", fontweight="bold", fontsizeasian="14pt", fontweightasian="bold"))
    o.styles.addElement(h2)
    mk = Style(name="Mark", family="paragraph")
    mk.addElement(TextProperties(color="#c83c32", fontsize="9pt", fontsizeasian="9pt"))
    mk.addElement(ParagraphProperties(textalign="end"))
    o.automaticstyles.addElement(mk)
    if doc.get("letterhead"):
        o.text.addElement(P(stylename=mk, text=doc["letterhead"]))
    o.text.addElement(H(outlinelevel=1, stylename=h1, text=doc["title"]))
    sep = "：" if doc.get("lang", "zh") == "zh" else ": "
    for b in doc["blocks"]:
        k = b[0]
        if k == "h":
            o.text.addElement(H(outlinelevel=2, stylename=h2, text=b[1]))
        elif k == "p":
            o.text.addElement(P(text=b[1]))
        elif k == "kv":
            for a, v in b[1]:
                p = P()
                p.addElement(Span(stylename=bold, text=f"{a}{sep}"))
                p.addText(str(v))
                o.text.addElement(p)
        elif k == "ul":
            lst = List()
            for x in b[1]:
                li = ListItem()
                li.addElement(P(text=x))
                lst.addElement(li)
            o.text.addElement(lst)
        elif k == "table":
            t = b[1]
            if t.get("caption"):
                cp = P()
                cp.addElement(Span(stylename=bold, text=t["caption"]))
                o.text.addElement(cp)
            tb = Table(name=f"T{id(t) % 10000}")
            tb.addElement(TableColumn(numbercolumnsrepeated=len(t["columns"])))
            for ri, r in enumerate([t["columns"]] + t["rows"]):
                tr = TableRow()
                for c in r:
                    tc = TableCell()
                    tc.addElement(P(text=cell_str(c)))
                    tr.addElement(tc)
                tb.addElement(tr)
            o.text.addElement(tb)
    if doc.get("signature"):
        o.text.addElement(P(text=doc["signature"]))
    o.text.addElement(P(stylename=mk, text=mark_for(doc)))
    o.save(path)


# ----------------------------------------------------------------------------- spreadsheets
# sheet spec: {"name", "title": merged title row across all columns (optional),
#              "group_header": [(c0, c1, text)] merged group labels in a row above the column header (optional),
#              "columns": [...], "rows": [[value | {"f": "=SUM(B3:B5)", "v": cached value}]],
#              "vmerge": [(col, r0, r1)] data rows r0..r1 of col merged (value from r0), "formats": {col: num|dec|pct},
#              "col_widths": [...]}
# Formulas use A1 references of the written layout; header_rows(sheet) tells where data starts.

def header_rows(sheet: dict) -> int:
    return (1 if sheet.get("title") else 0) + (1 if sheet.get("group_header") else 0) + 1


def data_ref(sheet: dict, col: int, i: int) -> str:
    """A1 reference of data row i (0-based), column col (0-based)."""
    return f"{_col(col)}{header_rows(sheet) + i + 1}"


def _col(c: int) -> str:
    s = ""
    c += 1
    while c:
        c, rem = divmod(c - 1, 26)
        s = chr(65 + rem) + s
    return s


def write_xlsx(path: str, sheets: list, props: dict | None = None):
    import xlsxwriter
    wb = xlsxwriter.Workbook(path, {"strings_to_numbers": False})
    wb.set_properties({"title": (props or {}).get("title", ""), "author": "synthetic", "comments": MARK,
                       "created": FIXED_TIME})
    hdr = wb.add_format({"bold": True, "bg_color": "#DDE6F2", "border": 1, "align": "center", "valign": "vcenter"})
    ttl = wb.add_format({"bold": True, "font_size": 14, "align": "center", "valign": "vcenter", "border": 1})
    fmts = {None: wb.add_format({"border": 1, "valign": "vcenter"}),
            "num": wb.add_format({"border": 1, "num_format": "#,##0"}),
            "dec": wb.add_format({"border": 1, "num_format": "0.00"}),
            "dec1": wb.add_format({"border": 1, "num_format": "0.0"}),
            "pct": wb.add_format({"border": 1, "num_format": "0.0%"})}
    markf = wb.add_format({"font_color": "#C83C32", "font_size": 9})
    for sh in sheets:
        ws = wb.add_worksheet(sh["name"])
        ncol = len(sh["columns"])
        r = 0
        if sh.get("title"):
            ws.merge_range(r, 0, r, ncol - 1, sh["title"], ttl)
            ws.set_row(r, 26)
            r += 1
        if sh.get("group_header"):
            for (c0, c1, text) in sh["group_header"]:
                if c1 > c0:
                    ws.merge_range(r, c0, r, c1, text, hdr)
                else:
                    ws.write_string(r, c0, text, hdr)
            r += 1
        for ci, c in enumerate(sh["columns"]):
            ws.write_string(r, ci, c, hdr)
        r += 1
        vcovered = {}
        for (col, a, b) in sh.get("vmerge", []):
            vcovered[(a, col)] = b
            for k in range(a + 1, b + 1):
                vcovered[(k, col)] = None
        for ri, row in enumerate(sh["rows"]):
            for ci, v in enumerate(row):
                fmt = fmts.get((sh.get("formats") or {}).get(ci), fmts[None])
                if (ri, ci) in vcovered:
                    end = vcovered[(ri, ci)]
                    if end is not None:
                        ws.merge_range(r + ri, ci, r + end, ci, v, fmt)
                    continue
                if isinstance(v, dict) and "f" in v:
                    ws.write_formula(r + ri, ci, v["f"], fmt, v["v"])
                elif isinstance(v, (int, float)) and not isinstance(v, bool):
                    ws.write_number(r + ri, ci, v, fmt)
                elif v is None:
                    ws.write_blank(r + ri, ci, None, fmt)
                else:
                    ws.write_string(r + ri, ci, str(v), fmt)
        for ci, wdt in enumerate(sh.get("col_widths", [])):
            ws.set_column(ci, ci, wdt)
        ws.write_string(r + len(sh["rows"]) + 1, 0, MARK, markf)
    wb.close()


def a1_to_odf(formula: str) -> str:
    """'=SUM(B2:B4)' -> 'of:=SUM([.B2:.B4])'; "='比价'!B5" -> "of:=[$'比价'.B5]"."""
    f = formula.lstrip("=")
    out, i = [], 0
    pat = re.compile(r"(?:'([^']+)'|([A-Za-z\u4e00-\u9fff_][\w\u4e00-\u9fff]*))!(\$?[A-Z]{1,2}\$?\d+)(?::(\$?[A-Z]{1,2}\$?\d+))?"
                     r"|(\$?[A-Z]{1,2}\$?\d+)(?::(\$?[A-Z]{1,2}\$?\d+))?")
    for m in pat.finditer(f):
        out.append(f[i:m.start()])
        if m.group(3):
            sh = m.group(1) or m.group(2)
            ref = f"[$'{sh}'.{m.group(3)}" + (f":.{m.group(4)}]" if m.group(4) else "]")
        else:
            ref = f"[.{m.group(5)}" + (f":.{m.group(6)}]" if m.group(6) else "]")
        out.append(ref)
        i = m.end()
    out.append(f[i:])
    return "of:=" + "".join(out).replace(",", ";")


def write_ods(path: str, sheets: list, props: dict | None = None):
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import Table, TableRow, TableCell, CoveredTableCell, TableColumn
    from odf.text import P
    from odf.style import Style, TextProperties, TableCellProperties
    from odf import dc, meta
    o = OpenDocumentSpreadsheet()
    o.meta.addElement(dc.Title(text=(props or {}).get("title", "")))
    o.meta.addElement(meta.CreationDate(text=FIXED_TIME.isoformat()))
    hs = Style(name="hdr", family="table-cell")
    hs.addElement(TextProperties(fontweight="bold", fontweightasian="bold"))
    hs.addElement(TableCellProperties(backgroundcolor="#DDE6F2", border="0.5pt solid #000000"))
    o.automaticstyles.addElement(hs)

    def value_cell(v, **kw):
        if isinstance(v, dict) and "f" in v:
            cv = v["v"]
            if isinstance(cv, (int, float)) and not isinstance(cv, bool):
                tc = TableCell(formula=a1_to_odf(v["f"]), valuetype="float", value=cv, **kw)
                tc.addElement(P(text=f"{cv:g}" if isinstance(cv, float) else str(cv)))
            else:
                tc = TableCell(formula=a1_to_odf(v["f"]), valuetype="string", stringvalue=str(cv), **kw)
                tc.addElement(P(text=str(cv)))
            return tc
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            tc = TableCell(valuetype="float", value=v, **kw)
            tc.addElement(P(text=f"{v:g}" if isinstance(v, float) else str(v)))
            return tc
        if v is None:
            return TableCell(**kw)
        tc = TableCell(valuetype="string", **kw)
        tc.addElement(P(text=str(v)))
        return tc

    for sh in sheets:
        ncol = len(sh["columns"])
        t = Table(name=sh["name"])
        t.addElement(TableColumn(numbercolumnsrepeated=ncol))
        if sh.get("title"):
            tr = TableRow()
            tc = TableCell(stylename=hs, valuetype="string", numbercolumnsspanned=ncol, numberrowsspanned=1)
            tc.addElement(P(text=sh["title"]))
            tr.addElement(tc)
            for _ in range(ncol - 1):
                tr.addElement(CoveredTableCell())
            t.addElement(tr)
        if sh.get("group_header"):
            tr = TableRow()
            spans = {c0: (c1, text) for c0, c1, text in sh["group_header"]}
            ci = 0
            while ci < ncol:
                if ci in spans:
                    c1, text = spans[ci]
                    tc = TableCell(stylename=hs, valuetype="string", numbercolumnsspanned=c1 - ci + 1, numberrowsspanned=1)
                    tc.addElement(P(text=text))
                    tr.addElement(tc)
                    for _ in range(c1 - ci):
                        tr.addElement(CoveredTableCell())
                    ci = c1 + 1
                else:
                    tr.addElement(TableCell())
                    ci += 1
            t.addElement(tr)
        tr = TableRow()
        for c in sh["columns"]:
            tc = TableCell(stylename=hs, valuetype="string")
            tc.addElement(P(text=c))
            tr.addElement(tc)
        t.addElement(tr)
        vstart = {(a, col): b - a + 1 for (col, a, b) in sh.get("vmerge", [])}
        vcov = {(k, col) for (col, a, b) in sh.get("vmerge", []) for k in range(a + 1, b + 1)}
        for ri, row in enumerate(sh["rows"]):
            tr = TableRow()
            for ci, v in enumerate(row):
                if (ri, ci) in vcov:
                    tr.addElement(CoveredTableCell())
                elif (ri, ci) in vstart:
                    tr.addElement(value_cell(v, numberrowsspanned=vstart[(ri, ci)], numbercolumnsspanned=1))
                else:
                    tr.addElement(value_cell(v))
            t.addElement(tr)
        t.addElement(TableRow())
        tr = TableRow()
        tc = TableCell(valuetype="string")
        tc.addElement(P(text=MARK))
        tr.addElement(tc)
        t.addElement(tr)
        o.spreadsheet.addElement(t)
    o.save(path)


def write_csv(path: str, columns: list[str], rows: list[list], encoding: str = "utf-8", delimiter: str = ",",
              newline: str = "\r\n", comment: str | None = None):
    import csv
    buf = io.StringIO()
    wr = csv.writer(buf, delimiter=delimiter, lineterminator=newline, quoting=csv.QUOTE_MINIMAL)
    if comment:
        wr.writerow([comment])
    wr.writerow(columns)
    for r in rows:
        wr.writerow([c if not isinstance(c, float) else f"{c:g}" for c in r])
    with open(path, "wb") as fh:
        fh.write(buf.getvalue().encode(encoding))


# ----------------------------------------------------------------------------- slides

def slide_picture(s: dict, theme: int = 0) -> Image.Image:
    """Picture of a slide (for picture-only slides and video/GIF frames)."""
    if s.get("chart"):
        c = s["chart"]
        return render_chart(c["title"], c["series"], c["xlabels"], kind=c.get("kind", "bar"), w=1280, h=720)
    bullets = list(s.get("bullets", []))
    if s.get("subtitle"):
        bullets = [s["subtitle"]] + bullets
    return render_slide_image(s["title"], bullets, theme=theme, table=s.get("table"), big=s.get("big"))


def write_pptx(path: str, slides: list[dict], title: str, tmpdir: str, theme: int = 0):
    """slides: {"title", "bullets", "image": bool (image-only slide), "notes", "table", "big"}."""
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    cp = prs.core_properties
    cp.title, cp.author, cp.created, cp.modified = title, "synthetic", FIXED_TIME, FIXED_TIME
    for i, s in enumerate(slides):
        if s.get("image"):
            sl = prs.slides.add_slide(prs.slide_layouts[6])
            im = slide_picture(s, theme)
            p = os.path.join(tmpdir, f"slide{i}.png")
            im.save(p)
            sl.shapes.add_picture(p, 0, 0, width=prs.slide_width, height=prs.slide_height)
        else:
            layout = prs.slide_layouts[0] if (i == 0 and not s.get("bullets")) else prs.slide_layouts[1]
            sl = prs.slides.add_slide(layout)
            sl.shapes.title.text = s["title"]
            if layout == prs.slide_layouts[0]:
                if len(sl.placeholders) > 1:
                    sl.placeholders[1].text = s.get("subtitle", "")
            else:
                body = sl.placeholders[1].text_frame
                bl = s.get("bullets", [])
                body.text = bl[0] if bl else ""
                for b in bl[1:]:
                    body.add_paragraph().text = b
                if s.get("table"):
                    t = s["table"]
                    sl.placeholders[1].height = Inches(2.2)
                    rows, cols = len(t["rows"]) + 1, len(t["columns"])
                    gt = sl.shapes.add_table(rows, cols, Inches(0.8), Inches(3.9), Inches(11.5), Inches(0.4) * rows).table
                    for ci, c in enumerate(t["columns"]):
                        gt.cell(0, ci).text = c
                    for ri, r in enumerate(t["rows"]):
                        for ci, c in enumerate(r):
                            gt.cell(ri + 1, ci).text = cell_str(c)
            tb = sl.shapes.add_textbox(Inches(11.3), Inches(7.0), Inches(1.9), Inches(0.4))
            tb.text_frame.text = MARK
            tb.text_frame.paragraphs[0].runs[0].font.size = Pt(10)
            tb.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xC8, 0x3C, 0x32)
        if s.get("notes"):
            sl.notes_slide.notes_text_frame.text = s["notes"]
    prs.save(path)


def write_odp(path: str, slides: list[dict], title: str, tmpdir: str, theme: int = 2):
    from odf.opendocument import OpenDocumentPresentation
    from odf.style import (Style, MasterPage, PageLayout, PageLayoutProperties, TextProperties, GraphicProperties,
                           ParagraphProperties, DrawingPageProperties)
    from odf.text import P, List, ListItem
    from odf.draw import Page, Frame, TextBox, Image as DImage
    from odf import dc, meta
    o = OpenDocumentPresentation()
    o.meta.addElement(dc.Title(text=title))
    o.meta.addElement(meta.CreationDate(text=FIXED_TIME.isoformat()))
    pl = PageLayout(name="PL")
    o.automaticstyles.addElement(pl)
    pl.addElement(PageLayoutProperties(margin="0cm", pagewidth="28cm", pageheight="15.75cm", printorientation="landscape"))
    mp = MasterPage(name="Master", pagelayoutname=pl)
    o.masterstyles.addElement(mp)
    dps = Style(name="dp1", family="drawing-page")
    dps.addElement(DrawingPageProperties(backgroundvisible="true"))
    o.automaticstyles.addElement(dps)
    ts = Style(name="title", family="presentation")
    ts.addElement(ParagraphProperties(textalign="start"))
    ts.addElement(TextProperties(fontsize="30pt", fontweight="bold", fontsizeasian="30pt", fontweightasian="bold",
                                 color="#1f3b63"))
    ts.addElement(GraphicProperties(fillcolor="#ffffff", stroke="none"))
    o.styles.addElement(ts)
    bs = Style(name="body", family="presentation")
    bs.addElement(TextProperties(fontsize="20pt", fontsizeasian="20pt"))
    bs.addElement(GraphicProperties(fill="none", stroke="none"))
    o.styles.addElement(bs)
    ms = Style(name="mark", family="presentation")
    ms.addElement(TextProperties(fontsize="10pt", fontsizeasian="10pt", color="#c83c32"))
    ms.addElement(GraphicProperties(fill="none", stroke="none"))
    o.styles.addElement(ms)
    ps = Style(name="photo", family="graphic")
    ps.addElement(GraphicProperties(fill="none", stroke="none"))
    o.automaticstyles.addElement(ps)
    for i, s in enumerate(slides):
        page = Page(stylename=dps, masterpagename=mp, name=f"page{i + 1}")
        o.presentation.addElement(page)
        if s.get("image"):
            im = slide_picture(s, theme)
            p = os.path.join(tmpdir, f"odp{i}.png")
            im.save(p)
            href = o.addPicture(p)
            fr = Frame(stylename=ps, width="28cm", height="15.75cm", x="0cm", y="0cm")
            page.addElement(fr)
            fr.addElement(DImage(href=href))
            continue
        fr = Frame(stylename=ts, width="25cm", height="2.4cm", x="1.5cm", y="0.8cm")
        page.addElement(fr)
        tb = TextBox()
        fr.addElement(tb)
        tb.addElement(P(text=s["title"]))
        lines = list(s.get("bullets", []))
        if s.get("subtitle"):
            lines = [s["subtitle"]] + lines
        if s.get("table"):
            t = s["table"]
            lines += ["  ".join(t["columns"])] + ["  ".join(cell_str(c) for c in r) for r in t["rows"]]
        fr2 = Frame(stylename=bs, width="25cm", height="11cm", x="1.5cm", y="3.6cm")
        page.addElement(fr2)
        tb2 = TextBox()
        fr2.addElement(tb2)
        lst = List()
        for ln in lines:
            li = ListItem()
            li.addElement(P(text=ln))
            lst.addElement(li)
        tb2.addElement(lst)
        fr3 = Frame(stylename=ms, width="4cm", height="0.8cm", x="23.5cm", y="14.8cm")
        page.addElement(fr3)
        tb3 = TextBox()
        fr3.addElement(tb3)
        tb3.addElement(P(text=MARK))
    o.save(path)


# ----------------------------------------------------------------------------- PDF

def write_pdf_text(path: str, doc: dict, pages_hint: int | None = None):
    """PDF with a real text layer (embedded subset TrueType font)."""
    import fitz
    pdf = fitz.open()
    W, H, M = 595, 842, 56
    fz = {"title": 20, "h": 14, "p": 10.5, "small": 8.5}
    fnt = fitz.Font(fontfile=CJK_TTF)
    state = {"page": None, "y": 0}

    def newpage():
        pg = pdf.new_page(width=W, height=H)
        pg.insert_font(fontname="cjk", fontfile=CJK_TTF)
        if doc.get("letterhead"):
            pg.insert_text((M, 36), doc["letterhead"], fontname="cjk", fontsize=fz["small"], color=(0.45, 0.45, 0.45))
            pg.draw_line((M, 42), (W - M, 42), color=(0.7, 0.7, 0.7), width=0.6)
        pg.insert_text((W - M - fnt.text_length(MARK, fz["small"]), H - 28), MARK, fontname="cjk",
                       fontsize=fz["small"], color=(0.78, 0.24, 0.2))
        pg.insert_text((M, H - 28), f"{len(pdf)}", fontname="cjk", fontsize=fz["small"], color=(0.5, 0.5, 0.5))
        state["page"], state["y"] = pg, M + 10

    def need(h):
        if state["page"] is None or state["y"] + h > H - M:
            newpage()

    def wrap_pdf(text, size, width):
        out, cur = [], ""
        for tok in re.findall(r"[A-Za-z0-9_\-./:@%#+()'\",$]+|\s+|.", text):
            cand = cur + tok
            if fnt.text_length(cand, size) <= width or not cur:
                cur = cand
            else:
                out.append(cur.rstrip())
                cur = tok.lstrip()
        if cur.strip():
            out.append(cur.rstrip())
        return out or [""]

    def line(text, size, x=None, color=(0, 0, 0), center=False):
        need(size * 1.6)
        pg = state["page"]
        xx = (W - fnt.text_length(text, size)) / 2 if center else (x if x is not None else M)
        pg.insert_text((xx, state["y"] + size), text, fontname="cjk", fontsize=size, color=color)
        state["y"] += size * 1.6

    newpage()
    for t in wrap_pdf(doc["title"], fz["title"], W - 2 * M):
        line(t, fz["title"], center=True)
    if doc.get("subtitle"):
        for t in wrap_pdf(doc["subtitle"], fz["p"], W - 2 * M):
            line(t, fz["p"], center=True, color=(0.3, 0.3, 0.3))
    state["y"] += 8
    sep = "：" if doc.get("lang", "zh") == "zh" else ": "
    for b in doc["blocks"]:
        k = b[0]
        if k == "h":
            state["y"] += 6
            line(b[1], fz["h"])
        elif k == "p":
            for t in wrap_pdf(b[1], fz["p"], W - 2 * M):
                line(t, fz["p"])
            state["y"] += 4
        elif k == "kv":
            for a, v in b[1]:
                for t in wrap_pdf(f"{a}{sep}{v}", fz["p"], W - 2 * M):
                    line(t, fz["p"])
            state["y"] += 4
        elif k == "ul":
            for x in b[1]:
                for j, t in enumerate(wrap_pdf(x, fz["p"], W - 2 * M - 14)):
                    line(("• " if j == 0 else "  ") + t, fz["p"])
            state["y"] += 4
        elif k == "table":
            t = b[1]
            cols = t["columns"]
            rows = [[cell_str(c) for c in r] for r in t["rows"]]
            size = fz["small"] + 0.5
            widths = [max(fnt.text_length(str(c), size) for c in [cols[i]] + [r[i] for r in rows]) + 12 for i in range(len(cols))]
            if sum(widths) > W - 2 * M:
                size = size * (W - 2 * M) / sum(widths)
                widths = [max(fnt.text_length(str(c), size) for c in [cols[i]] + [r[i] for r in rows]) + 12 for i in range(len(cols))]
            rh = max(size, 7) * 2.1
            if t.get("caption"):
                line(t["caption"], fz["p"])
            for ri, r in enumerate([cols] + rows):
                need(rh)
                pg, y = state["page"], state["y"]
                x = M
                for ci, c in enumerate(r):
                    rect = fitz.Rect(x, y, x + widths[ci], y + rh)
                    pg.draw_rect(rect, color=(0.4, 0.4, 0.4), fill=(0.9, 0.92, 0.95) if ri == 0 else None, width=0.5)
                    pg.insert_text((x + 4, y + rh * 0.68), str(c), fontname="cjk", fontsize=size)
                    x += widths[ci]
                state["y"] += rh
            state["y"] += 8
    if doc.get("signature"):
        state["y"] += 16
        line(doc["signature"], fz["p"], x=W - M - 220)
    pdf.set_metadata({})  # no Info dictionary: nothing about the authoring tool or machine
    pdf.xref_set_key(-1, "Info", "null")
    pdf.xref_set_key(pdf.pdf_catalog(), "Info", "null")
    pdf.subset_fonts()
    pdf.save(path, garbage=4, deflate=True)
    pdf.close()


def write_pdf_scanned(path: str, doc: dict, rng: random.Random):
    """Image-only PDF: pages rendered, skewed, noised and JPEG-compressed; no text layer at all."""
    import fitz
    pages = render_page_images(doc)
    pdf = fitz.open()
    for pg in pages:
        im = scan_effect(pg, rng)
        buf = io.BytesIO()
        im.save(buf, "JPEG", quality=55, optimize=True)
        p = pdf.new_page(width=595, height=842)
        p.insert_image(p.rect, stream=buf.getvalue())
    pdf.set_metadata({})
    pdf.xref_set_key(-1, "Info", "null")
    pdf.xref_set_key(pdf.pdf_catalog(), "Info", "null")
    pdf.save(path, garbage=4, deflate=True)
    pdf.close()
    return len(pages)


# ----------------------------------------------------------------------------- EPUB / webarchive / zip

def _fixed_zipinfo(name: str, compress=zipfile.ZIP_DEFLATED) -> zipfile.ZipInfo:
    zi = zipfile.ZipInfo(name, date_time=(2026, 10, 1, 9, 0, 0))
    zi.compress_type = compress
    zi.external_attr = 0o644 << 16
    return zi


def write_epub(path: str, book: dict):
    """book: title, author, lang, identifier, chapters: [{"title", "blocks"}]."""
    e = html.escape
    lang = "zh-CN" if book.get("lang", "zh") == "zh" else "en"
    chap_files = []
    for i, ch in enumerate(book["chapters"], 1):
        parts = [f"<h1>{e(ch['title'])}</h1>"]
        for b in ch["blocks"]:
            if b[0] == "p":
                parts.append(f"<p>{e(b[1])}</p>")
            elif b[0] == "h":
                parts.append(f"<h2>{e(b[1])}</h2>")
            elif b[0] == "ul":
                parts.append("<ul>" + "".join(f"<li>{e(x)}</li>" for x in b[1]) + "</ul>")
            elif b[0] == "table":
                t = b[1]
                parts.append("<table><tr>" + "".join(f"<th>{e(c)}</th>" for c in t["columns"]) + "</tr>" +
                             "".join("<tr>" + "".join(f"<td>{e(cell_str(c))}</td>" for c in r) + "</tr>" for r in t["rows"])
                             + "</table>")
        if i == len(book["chapters"]):
            parts.append(f'<p class="mark">{e(MARK)}</p>')
        x = (f'<?xml version="1.0" encoding="utf-8"?>\n<!DOCTYPE html>\n<html xmlns="http://www.w3.org/1999/xhtml" '
             f'xmlns:epub="http://www.idpf.org/2007/ops" xml:lang="{lang}" lang="{lang}">\n<head><title>{e(ch["title"])}</title>'
             f'<link rel="stylesheet" type="text/css" href="style.css"/></head>\n<body>\n' + "\n".join(parts) + "\n</body>\n</html>\n")
        chap_files.append((f"chap{i}.xhtml", x, ch["title"]))
    nav = (f'<?xml version="1.0" encoding="utf-8"?>\n<!DOCTYPE html>\n<html xmlns="http://www.w3.org/1999/xhtml" '
           f'xmlns:epub="http://www.idpf.org/2007/ops" xml:lang="{lang}"><head><title>{e(book["title"])}</title></head><body>'
           f'<nav epub:type="toc" id="toc"><h1>{e(book["title"])}</h1><ol>' +
           "".join(f'<li><a href="{f}">{e(t)}</a></li>' for f, _, t in chap_files) + "</ol></nav></body></html>\n")
    manifest = "".join(f'<item id="c{i}" href="{f}" media-type="application/xhtml+xml"/>' for i, (f, _, _) in enumerate(chap_files, 1))
    spine = "".join(f'<itemref idref="c{i}"/>' for i in range(1, len(chap_files) + 1))
    opf = (f'<?xml version="1.0" encoding="utf-8"?>\n<package xmlns="http://www.idpf.org/2007/opf" version="3.0" '
           f'unique-identifier="bookid" xml:lang="{lang}">\n<metadata xmlns:dc="http://purl.org/dc/elements/1.1/">'
           f'<dc:identifier id="bookid">{e(book["identifier"])}</dc:identifier><dc:title>{e(book["title"])}</dc:title>'
           f'<dc:creator>{e(book["author"])}</dc:creator><dc:language>{lang}</dc:language>'
           f'<dc:description>{e(MARK)}</dc:description>'
           f'<meta property="dcterms:modified">2026-10-01T09:00:00Z</meta></metadata>\n'
           f'<manifest><item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>'
           f'<item id="css" href="style.css" media-type="text/css"/>{manifest}</manifest>\n<spine>{spine}</spine>\n</package>\n')
    container = ('<?xml version="1.0" encoding="utf-8"?>\n<container version="1.0" '
                 'xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles>'
                 '<rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>\n')
    css = "body{font-family:serif;line-height:1.6}.mark{color:#c83c32;font-size:small}table{border-collapse:collapse}td,th{border:1px solid #999;padding:2px 6px}"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr(_fixed_zipinfo("mimetype", zipfile.ZIP_STORED), "application/epub+zip")
        z.writestr(_fixed_zipinfo("META-INF/container.xml"), container)
        z.writestr(_fixed_zipinfo("OEBPS/content.opf"), opf)
        z.writestr(_fixed_zipinfo("OEBPS/nav.xhtml"), nav)
        z.writestr(_fixed_zipinfo("OEBPS/style.css"), css)
        for f, x, _ in chap_files:
            z.writestr(_fixed_zipinfo(f"OEBPS/{f}"), x)


def write_webarchive(path: str, url: str, html_text: str, subresources: list[tuple[str, str, bytes]]):
    """Safari .webarchive: binary plist with WebMainResource and WebSubresources; converted with plutil."""
    main = {"WebResourceData": html_text.encode("utf-8"), "WebResourceFrameName": "",
            "WebResourceMIMEType": "text/html", "WebResourceTextEncodingName": "UTF-8", "WebResourceURL": url}
    subs = [{"WebResourceData": data, "WebResourceMIMEType": mime, "WebResourceURL": u} for u, mime, data in subresources]
    pl = {"WebMainResource": main}
    if subs:
        pl["WebSubresources"] = subs
    tmp = path + ".xml"
    with open(tmp, "wb") as fh:
        plistlib.dump(pl, fh, fmt=plistlib.FMT_XML)
    subprocess.run(["/usr/bin/plutil", "-convert", "binary1", tmp, "-o", path], check=True)
    os.remove(tmp)
    subprocess.run(["/usr/bin/plutil", "-lint", "-s", path], check=True)


def write_zip(path: str, members: list[tuple[str, bytes]]):
    with zipfile.ZipFile(path, "w") as z:
        for name, data in members:
            z.writestr(_fixed_zipinfo(name), data)


# ----------------------------------------------------------------------------- images & video

def save_image(im: Image.Image, path: str, fmt: str, tmpdir: str):
    fmt = fmt.lower()
    if fmt == "png":
        im.save(path, "PNG", optimize=True)
    elif fmt == "jpg":
        im.convert("RGB").save(path, "JPEG", quality=82, optimize=True)
    elif fmt == "webp":
        im.save(path, "WEBP", quality=80, method=6)
    elif fmt == "bmp":
        (im if im.mode in ("P", "L", "1") else im.convert("RGB")).save(path, "BMP")
    elif fmt == "tiff":
        im.save(path, "TIFF", compression="tiff_lzw")
    elif fmt == "heic":
        src = os.path.join(tmpdir, "heic_src.jpg")
        im.convert("RGB").save(src, "JPEG", quality=92)
        subprocess.run(["/usr/bin/sips", "-s", "format", "heic", "-s", "formatOptions", "70", src, "--out", path],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    else:
        raise ValueError(fmt)


def save_tiff_pages(pages: list[Image.Image], path: str):
    pages[0].save(path, "TIFF", compression="tiff_lzw", save_all=True, append_images=pages[1:])


def save_gif(frames: list[Image.Image], path: str, durations: list[int]):
    pal = [f.convert("RGB").quantize(colors=128, method=Image.Quantize.MEDIANCUT, dither=Image.Dither.NONE) for f in frames]
    pal[0].save(path, "GIF", save_all=True, append_images=pal[1:], duration=durations, loop=0, disposal=2, optimize=False)


def ffmpeg_exe() -> str:
    import imageio_ffmpeg
    return imageio_ffmpeg.get_ffmpeg_exe()


def write_video(path: str, frames: list[Image.Image], seconds: list[float], tmpdir: str, silent_audio: bool = False,
                fps: int = 4):
    """Slides video: each frame held for its seconds; H.264 yuv420p; optional silent AAC track."""
    fdir = tempfile.mkdtemp(dir=tmpdir)
    n = 0
    for im, sec in zip(frames, seconds):
        im = im.convert("RGB")
        for _ in range(max(1, int(round(sec * fps)))):
            im.save(os.path.join(fdir, f"f{n:04d}.png"))
            n += 1
    cmd = [ffmpeg_exe(), "-y", "-loglevel", "error", "-framerate", str(fps), "-i", os.path.join(fdir, "f%04d.png")]
    if silent_audio:
        cmd += ["-f", "lavfi", "-i", "anullsrc=r=44100:cl=mono", "-shortest", "-c:a", "aac", "-b:a", "32k"]
    else:
        cmd += ["-an"]
    cmd += ["-c:v", "libx264", "-preset", "medium", "-crf", "30", "-pix_fmt", "yuv420p", "-tune", "stillimage",
            "-fflags", "+bitexact", "-flags:v", "+bitexact", "-map_metadata", "-1"]
    if path.endswith(".mp4"):
        cmd += ["-movflags", "+faststart"]
    cmd += [path]
    subprocess.run(cmd, check=True)
    shutil.rmtree(fdir, ignore_errors=True)
