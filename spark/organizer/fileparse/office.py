"""Office Open XML (docx / xlsx / pptx and their macro / template variants), OpenDocument (odt / ods / odp),
EPUB and Apple iWork packages (Pages / Numbers / Keynote via their embedded preview).

Only the XML parts are read (defusedxml). Macros (vbaProject.bin), OLE objects, external links and
ActiveX parts are never opened.
"""

from __future__ import annotations

import io
import posixpath
import re
from typing import Callable, Optional

from .core import (MAX_SHEET_CELLS_SCANNED, MAX_SHEET_COLS, MAX_SHEET_ROWS, MAX_SHEETS, Budget, ParseError, Parsed,
                   cell_text, clean_text, first_line, html_to_text, md_table, number_as_shown)
from .images import add_image
from .pkg import Package, local, rels

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"

IMAGE_MEMBER = re.compile(r"\.(png|jpe?g|gif|bmp|tiff?|webp)$", re.I)


def chart_text(pkg: Package, part: str) -> str:
    """A chart embedded in a docx / pptx: its title and the cached data of every series as a table
    (categories as rows). The workbook embedded with the chart is not opened."""
    try:
        root = pkg.xml(part)
    except ParseError:
        return ""
    if root is None:
        return ""
    title_el = root.find(f".//{C}title")
    title = "".join(t.text or "" for t in title_el.iter(f"{A}t")).strip() if title_el is not None else ""

    def points(el) -> dict[int, str]:
        out: dict[int, str] = {}
        if el is None:
            return out
        for pt in el.iter(f"{C}pt"):
            v = pt.find(f"{C}v")
            if v is not None and (pt.get("idx") or "").isdigit():
                out[int(pt.get("idx"))] = cell_text(_num(v.text or ""))
        return out

    series = []
    for ser in root.iter(f"{C}ser"):
        tx = ser.find(f"{C}tx")
        name = "".join(v.text or "" for v in tx.iter(f"{C}v")).strip() if tx is not None else ""
        cats = points(ser.find(f"{C}cat")) or points(ser.find(f"{C}xVal"))
        vals = points(ser.find(f"{C}val")) or points(ser.find(f"{C}yVal"))
        series.append((name or f"系列{len(series) + 1}", cats, vals))
    if not series:
        return f"图表：{title}" if title else ""
    idxs = sorted(set().union(*(set(c) | set(v) for _, c, v in series)))[:200]
    rows = [["类别"] + [name for name, _, _ in series]]
    for i in idxs:
        cat = next((c[i] for _, c, _ in series if i in c), str(i + 1))
        rows.append([cat] + [v.get(i, "") for _, _, v in series])
    table, _ = md_table(rows, max_rows=200)
    return (f"图表：{title}\n" if title else "图表：\n") + table


def _num(text: str):
    try:
        return float(text) if any(ch in text for ch in ".eE") else int(text)
    except ValueError:
        return text


def detect_package(pkg: Package, ext: str) -> str:
    """docx / xlsx / pptx / odt / ods / odp / epub / iwork / zip for a zip container."""
    try:
        mimetype = pkg.read("mimetype", limit=4096)
    except ParseError:  # an odd "mimetype" member does not decide the format
        mimetype = None
    if mimetype:
        fmt = {"application/vnd.oasis.opendocument.text": "odt",
               "application/vnd.oasis.opendocument.text-template": "odt",
               "application/vnd.oasis.opendocument.spreadsheet": "ods",
               "application/vnd.oasis.opendocument.spreadsheet-template": "ods",
               "application/vnd.oasis.opendocument.presentation": "odp",
               "application/vnd.oasis.opendocument.presentation-template": "odp",
               "application/epub+zip": "epub"}.get(mimetype.decode("ascii", "replace").strip())
        if fmt:
            return fmt
    for main, fmt in (("word/document.xml", "docx"), ("xl/workbook.xml", "xlsx"), ("ppt/presentation.xml", "pptx")):
        if pkg.has(main):
            return fmt
    if pkg.has("[Content_Types].xml"):
        try:
            ct = pkg.read("[Content_Types].xml", limit=1 << 20) or b""
        except ParseError:
            ct = b""
        for key, fmt in ((b"wordprocessingml", "docx"), (b"spreadsheetml", "xlsx"), (b"presentationml", "pptx")):
            if key in ct:
                return fmt
    names = [n.lower() for n in pkg.names()]
    if ext == "xmind" or ("content.json" in names and "metadata.json" in names) or \
            ("content.xml" in names and "meta.xml" in names and not pkg.has("mimetype")):
        return "xmind"
    if ext in ("pages", "numbers", "key") or any(n.startswith("index/") and n.endswith(".iwa") for n in names) \
            or "index.zip" in names:
        return "iwork"
    if pkg.has("META-INF/container.xml") and ext == "epub":
        return "epub"
    return "zip"


# ---- docx ----------------------------------------------------------------------------------------------


def _w_para(p, rel: dict, pkg: Package, budget: Budget, images: list) -> str:
    parts: list[str] = []
    for el in p.iter():
        tag = local(el.tag)
        if tag == "t" and el.tag == f"{W}t":
            parts.append(el.text or "")
        elif tag == "tab" and el.tag == f"{W}tab":
            parts.append("\t")
        elif tag in ("br", "cr") and el.tag.startswith(W):
            parts.append("\n")
        elif tag == "blip":
            rid = el.get(f"{R}embed")
            if rid and rid in rel:
                images.append(rel[rid])
                marker = _image_marker(pkg, rel[rid], budget, f"文档图片 {len(images)}")
                if marker:
                    parts.append(f"\n{marker}\n")
        elif el.tag == f"{C}chart":
            rid = el.get(f"{R}id")
            if rid and rid in rel:
                parts.append("\n" + chart_text(pkg, rel[rid]) + "\n")
    text = "".join(parts).strip()
    ppr = p.find(f"{W}pPr")
    if ppr is not None and text:
        style = ppr.find(f"{W}pStyle")
        sval = (style.get(f"{W}val") or "") if style is not None else ""
        m = re.search(r"(?i)(?:heading|标题)\s*(\d)", sval)
        if m or sval.lower() == "title":
            level = int(m.group(1)) if m else 1
            return "#" * min(level, 4) + " " + text
        if ppr.find(f"{W}numPr") is not None:
            return "- " + text
    return text


def _image_marker(pkg: Package, member: str, budget: Budget, label: str) -> Optional[str]:
    if not IMAGE_MEMBER.search(member):
        return None
    try:
        data = pkg.read(member, limit=40 * 1024 * 1024)
    except ParseError:
        return None
    return add_image(budget, data, label) if data else None


def _w_table(tbl, rel, pkg, budget, images) -> str:
    rows = []
    for tr in tbl.iter(f"{W}tr"):
        cells = []
        for tc in tr.findall(f"{W}tc"):
            paras = [_w_para(p, rel, pkg, budget, images) for p in tc.iter(f"{W}p")]
            cells.append(" / ".join(x for x in paras if x).replace("|", "\\|").replace("\n", " "))
        rows.append(cells)
    table, left = md_table(rows)
    return table + (f"\n（表格其余 {left} 行未列出）" if left else "")


def parse_docx(pkg: Package, budget: Budget) -> Parsed:
    root = pkg.xml("word/document.xml")
    if root is None:
        raise ParseError("corrupt", "docx without word/document.xml")
    rel = rels(pkg, "word/document.xml")
    body = root.find(f"{W}body")
    images: list[str] = []
    blocks: list[str] = []

    def walk(container) -> None:
        for child in container:
            tag = child.tag
            if tag == f"{W}p":
                t = _w_para(child, rel, pkg, budget, images)
                if t:
                    blocks.append(t)
            elif tag == f"{W}tbl":
                t = _w_table(child, rel, pkg, budget, images)
                if t:
                    blocks.append(t)
            elif tag in (f"{W}sdt", f"{W}sdtContent", f"{W}customXml", f"{W}ins", f"{W}smartTag"):
                walk(child)

    walk(body if body is not None else root)
    # Footnotes and comments are part of what the author wrote.
    for part, head in (("word/footnotes.xml", "脚注"), ("word/comments.xml", "批注")):
        try:
            extra = pkg.xml(part)
        except ParseError:
            extra = None
        if extra is not None:
            notes = [t for t in ("".join((x.text or "") for x in p.iter(f"{W}t")).strip()
                                 for p in extra.iter(f"{W}p")) if t]
            if notes:
                blocks.append(f"## {head}\n" + "\n".join(notes))
    text = clean_text("\n\n".join(blocks))
    title = _core_title(pkg) or first_line(text)
    return Parsed("document", text, title, {"images_found": len(images)})


def _core_title(pkg: Package) -> str:
    try:
        core = pkg.xml("docProps/core.xml")
    except ParseError:
        return ""
    if core is None:
        return ""
    for el in core:
        if local(el.tag) == "title" and (el.text or "").strip():
            return el.text.strip()[:120]
    return ""


# ---- pptx ----------------------------------------------------------------------------------------------


def _slide_text(root, rel: dict, pkg: Package, budget: Budget, slide_no: int) -> list[str]:
    lines: list[str] = []
    n_img = 0

    def shape_is_title(sp) -> bool:
        ph = sp.find(f".//{P}nvPr/{P}ph")
        return ph is not None and (ph.get("type") or "") in ("title", "ctrTitle")

    def walk(el) -> None:
        nonlocal n_img
        for child in el:
            tag = child.tag
            if tag == f"{P}sp":
                title = shape_is_title(child)
                for para in child.iter(f"{A}p"):
                    t = "".join((x.text or "") if x.tag == f"{A}t" else ("\n" if x.tag == f"{A}br" else "")
                                for x in para.iter()).strip()
                    if t:
                        lvl = para.find(f"{A}pPr")
                        indent = int(lvl.get("lvl", "0")) if lvl is not None and (lvl.get("lvl") or "0").isdigit() else 0
                        lines.append(("### " + t) if title else ("  " * indent + "- " + t))
            elif tag == f"{A}tbl" or local(tag) == "graphicFrame":
                tbl = child if tag == f"{A}tbl" else child.find(f".//{A}tbl")
                if tbl is not None:
                    rows = [[" ".join((t.text or "") for t in tc.iter(f"{A}t")).strip() for tc in tr.findall(f"{A}tc")]
                            for tr in tbl.iter(f"{A}tr")]
                    table, _ = md_table(rows, max_rows=60)
                    if table:
                        lines.append(table)
                else:
                    walk(child)
            elif tag == f"{C}chart":
                rid = child.get(f"{R}id")
                if rid and rid in rel:
                    t = chart_text(pkg, rel[rid])
                    if t:
                        lines.append(t)
            elif tag == f"{P}pic":
                blip = child.find(f".//{A}blip")
                rid = blip.get(f"{R}embed") if blip is not None else None
                if rid and rid in rel:
                    n_img += 1
                    marker = _image_marker(pkg, rel[rid], budget, f"第 {slide_no} 张幻灯片的图片 {n_img}")
                    if marker:
                        lines.append(marker)
            else:
                walk(child)

    walk(root)
    return lines


def parse_pptx(pkg: Package, budget: Budget) -> Parsed:
    pres = pkg.xml("ppt/presentation.xml")
    if pres is None:
        raise ParseError("corrupt", "pptx without ppt/presentation.xml")
    prel = rels(pkg, "ppt/presentation.xml")
    order = []
    lst = pres.find(f"{P}sldIdLst")
    if lst is not None:
        for sid in lst:
            rid = sid.get(f"{R}id")
            if rid in prel:
                order.append(prel[rid])
    if not order:
        order = sorted((n for n in pkg.names() if re.match(r"ppt/slides/slide\d+\.xml$", n)),
                       key=lambda n: int(re.search(r"(\d+)", n.rsplit("/", 1)[-1]).group(1)))
    blocks: list[str] = []
    title = ""
    for i, part in enumerate(order, 1):
        try:
            root = pkg.xml(part)
        except ParseError:
            root = None
        if root is None:
            continue
        srel = rels(pkg, part)
        lines = _slide_text(root, srel, pkg, budget, i)
        notes_part = next((t for t in srel.values() if "notesSlide" in t), None)
        if notes_part:
            try:
                nroot = pkg.xml(notes_part)
            except ParseError:
                nroot = None
            if nroot is not None:
                note = " ".join(t for t in ("".join((x.text or "") for x in p.iter(f"{A}t")).strip()
                                            for p in nroot.iter(f"{A}p")) if t and not t.isdigit())
                if note:
                    lines.append(f"备注：{note}")
        head = next((ln[4:] for ln in lines if ln.startswith("### ")), "")
        if i == 1 and head:
            title = head
        blocks.append(f"## 第 {i} 张幻灯片" + (f"：{head}" if head else "") + "\n"
                      + "\n".join(ln for ln in lines if not (ln.startswith("### ") and ln[4:] == head)))
    text = clean_text("\n\n".join(blocks))
    return Parsed("slides", text, _core_title(pkg) or title or first_line(text), {"slides": len(order)})


# ---- xlsx ----------------------------------------------------------------------------------------------


def parse_xlsx(data: bytes, pkg: Package, budget: Budget) -> Parsed:
    # Declared sizes are checked first (Package.read raises on a bomb), then openpyxl streams the sheets.
    for name in pkg.names():
        if name.startswith("xl/"):
            info = pkg.infos[name]
            from .pkg import is_bomb
            if is_bomb(info) or info.file_size > 400 * 1024 * 1024:
                raise ParseError("too_large", f"{name}: {info.file_size} bytes uncompressed")
    import openpyxl  # MIT; uses defusedxml when it is installed
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    except Exception as exc:  # noqa: BLE001
        raise ParseError("corrupt", f"xlsx: {type(exc).__name__}: {exc}") from exc
    blocks: list[str] = []
    scanned = 0
    sheets = wb.sheetnames
    for si, name in enumerate(sheets):
        if si >= MAX_SHEETS:
            budget.note(f"工作表超过 {MAX_SHEETS} 个，其余未读")
            break
        ws = wb[name]
        rows: list[list[str]] = []
        total = 0
        try:
            for row in ws.iter_rows(max_col=MAX_SHEET_COLS):
                total += 1
                scanned += len(row)
                if len(rows) <= MAX_SHEET_ROWS:
                    # As displayed: percentages, fixed decimals, thousands separators, currency signs.
                    rows.append([number_as_shown(getattr(c, "value", None), getattr(c, "number_format", "") or "")
                                 for c in row])
                if scanned > MAX_SHEET_CELLS_SCANNED:
                    budget.note("表格过大，只读了前面的部分")
                    break
        except Exception as exc:  # noqa: BLE001
            raise ParseError("corrupt", f"sheet {name}: {type(exc).__name__}") from exc
        table, left = md_table(rows)
        non_empty = sum(1 for r in rows if any(c for c in r))
        more = max(0, total - len(rows))
        hidden = f"（其余 {left + more} 行未列出）" if left + more else ""
        dims = getattr(ws, "max_row", None)
        blocks.append(f"## 工作表：{name}" + (f"（{dims} 行）" if dims and dims > non_empty else "") + "\n"
                      + (table or "（空表）") + (f"\n{hidden}" if hidden else ""))
        if scanned > MAX_SHEET_CELLS_SCANNED:
            break
    wb.close()
    text = "\n\n".join(blocks)
    return Parsed("spreadsheet", text, "", {"sheets": len(sheets)})


# ---- OpenDocument ----------------------------------------------------------------------------------------

TEXT_NS = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"
TABLE_NS = "{urn:oasis:names:tc:opendocument:xmlns:table:1.0}"
DRAW_NS = "{urn:oasis:names:tc:opendocument:xmlns:drawing:1.0}"
XLINK = "{http://www.w3.org/1999/xlink}"
OFFICE_NS = "{urn:oasis:names:tc:opendocument:xmlns:office:1.0}"


def _odf_encrypted(pkg: Package) -> bool:
    manifest = pkg.read("META-INF/manifest.xml", limit=4 << 20) or b""
    return b"encryption-data" in manifest


def _odf_text(el) -> str:
    out: list[str] = []

    def rec(e) -> None:
        if e.text:
            out.append(e.text)
        for c in e:
            tag = c.tag
            if tag == f"{TEXT_NS}s":
                out.append(" " * int(c.get(f"{TEXT_NS}c", "1") or 1) if (c.get(f"{TEXT_NS}c", "1") or "1").isdigit() else " ")
            elif tag == f"{TEXT_NS}tab":
                out.append("\t")
            elif tag == f"{TEXT_NS}line-break":
                out.append("\n")
            elif tag in (f"{TEXT_NS}note", f"{OFFICE_NS}annotation"):
                pass
            else:
                rec(c)
            if c.tail:
                out.append(c.tail)

    rec(el)
    return "".join(out).strip()


def _odf_rows(table, max_rows: int = MAX_SHEET_ROWS) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    total = 0
    for row in table.iter(f"{TABLE_NS}table-row"):
        rep = int(row.get(f"{TABLE_NS}number-rows-repeated", "1") or 1)
        cells: list[str] = []
        for cell in row:
            if local(cell.tag) not in ("table-cell", "covered-table-cell"):
                continue
            crep = min(int(cell.get(f"{TABLE_NS}number-columns-repeated", "1") or 1), MAX_SHEET_COLS)
            value = " / ".join(t for t in (_odf_text(p) for p in cell.iter(f"{TEXT_NS}p")) if t)
            cells.extend([cell_text(value)] * crep)
            if len(cells) >= MAX_SHEET_COLS:
                break
        if not any(cells):
            total += min(rep, 1)
            continue  # repeated empty rows pad a sheet to a million rows: never expanded
        for _ in range(min(rep, max_rows + 1)):
            total += 1
            if len(rows) <= max_rows:
                rows.append(cells[:MAX_SHEET_COLS])
    return rows, total


def parse_odf(pkg: Package, fmt: str, budget: Budget) -> Parsed:
    if _odf_encrypted(pkg):
        raise ParseError("encrypted", "OpenDocument with encryption-data")
    root = pkg.xml("content.xml")
    if root is None:
        raise ParseError("corrupt", "OpenDocument without content.xml")
    body = root.find(f"{OFFICE_NS}body")
    body = body if body is not None else root
    blocks: list[str] = []
    counts: dict = {}
    if fmt == "ods":
        tables = list(body.iter(f"{TABLE_NS}table"))
        for t in tables[:MAX_SHEETS]:
            rows, total = _odf_rows(t)
            table, left = md_table(rows)
            more = max(0, total - len(rows)) + left
            blocks.append(f"## 工作表：{t.get(f'{TABLE_NS}name', '')}\n" + (table or "（空表）")
                          + (f"\n（其余 {more} 行未列出）" if more else ""))
        counts["sheets"] = len(tables)
        typ = "spreadsheet"
    elif fmt == "odp":
        pages = list(body.iter(f"{DRAW_NS}page"))
        for i, page in enumerate(pages, 1):
            lines = [t for t in (_odf_text(p) for p in page.iter(f"{TEXT_NS}p")) if t]
            imgs = []
            for img in page.iter(f"{DRAW_NS}image"):
                href = img.get(f"{XLINK}href") or ""
                marker = _image_marker(pkg, href, budget, f"第 {i} 张幻灯片的图片") if href else None
                if marker:
                    imgs.append(marker)
            blocks.append(f"## 第 {i} 张幻灯片\n" + "\n".join(lines + imgs))
        counts["slides"] = len(pages)
        typ = "slides"
    else:
        def walk(el) -> None:
            for c in el:
                tag = c.tag
                if tag == f"{TEXT_NS}h":
                    t = _odf_text(c)
                    if t:
                        blocks.append("#" * min(int(c.get(f"{TEXT_NS}outline-level", "1") or 1), 4) + " " + t)
                elif tag == f"{TEXT_NS}p":
                    t = _odf_text(c)
                    for img in c.iter(f"{DRAW_NS}image"):
                        href = img.get(f"{XLINK}href") or ""
                        marker = _image_marker(pkg, href, budget, "文档图片") if href else None
                        if marker:
                            t += f"\n{marker}"
                    if t:
                        blocks.append(t)
                elif tag == f"{TEXT_NS}list":
                    for item in c.iter(f"{TEXT_NS}list-item"):
                        t = " ".join(x for x in (_odf_text(p) for p in item.findall(f"{TEXT_NS}p")) if x)
                        if t:
                            blocks.append("- " + t)
                elif tag == f"{TABLE_NS}table":
                    rows, _ = _odf_rows(c, 200)
                    table, _ = md_table(rows)
                    if table:
                        blocks.append(table)
                else:
                    walk(c)
        walk(body)
        typ = "document"
    text = clean_text("\n\n".join(blocks))
    return Parsed(typ, text, first_line(text), counts)


# ---- EPUB ------------------------------------------------------------------------------------------------


def parse_epub(pkg: Package, budget: Budget) -> Parsed:
    if pkg.has("META-INF/encryption.xml"):
        enc = pkg.read("META-INF/encryption.xml", limit=4 << 20) or b""
        if b"EncryptedData" in enc and b"obfuscation" not in enc.lower():
            raise ParseError("encrypted", "EPUB with DRM (META-INF/encryption.xml)")
    container = pkg.xml("META-INF/container.xml")
    opf_path = None
    if container is not None:
        for el in container.iter():
            if local(el.tag) == "rootfile" and el.get("full-path"):
                opf_path = el.get("full-path")
                break
    if not opf_path:
        opf_path = next((n for n in pkg.names() if n.lower().endswith(".opf")), None)
    if not opf_path:
        raise ParseError("corrupt", "EPUB without a package document")
    opf = pkg.xml(opf_path)
    folder = posixpath.dirname(opf_path)
    manifest, spine, meta = {}, [], {}
    for el in opf.iter():
        tag = local(el.tag)
        if tag == "item" and el.get("id"):
            manifest[el.get("id")] = (posixpath.normpath(posixpath.join(folder, el.get("href") or "")),
                                      el.get("media-type") or "")
        elif tag == "itemref" and el.get("idref"):
            spine.append(el.get("idref"))
        elif tag in ("title", "creator", "language", "publisher", "date") and (el.text or "").strip():
            meta.setdefault(tag, el.text.strip())
    blocks = []
    head = []
    if meta.get("title"):
        head.append(f"书名：{meta['title']}")
    if meta.get("creator"):
        head.append(f"作者：{meta['creator']}")
    if meta.get("publisher"):
        head.append(f"出版：{meta['publisher']}")
    if head:
        blocks.append("\n".join(head))
    chapters = 0
    size = 0
    for idref in spine:
        path, mtype = manifest.get(idref, ("", ""))
        if not path or "html" not in mtype and not path.lower().endswith((".xhtml", ".html", ".htm")):
            continue
        try:
            raw = pkg.read(path, limit=20 * 1024 * 1024)
        except ParseError:
            continue
        if not raw:
            continue
        from .core import decode_text
        _, text = html_to_text(decode_text(raw))
        if text:
            chapters += 1
            blocks.append(text)
            size += len(text)
            if size > 200_000:
                budget.note("电子书较长，只读了前面的章节")
                break
    text = clean_text("\n\n".join(blocks))
    fields = [{"key": k, "label": lab, "value": meta[k]} for k, lab in (("title", "书名"), ("creator", "作者"))
              if meta.get(k)]
    return Parsed("ebook", text, meta.get("title", "") or first_line(text), {"chapters": len(spine)}, fields=fields)


# ---- iWork -----------------------------------------------------------------------------------------------


def parse_iwork(pkg: Package, ext: str, budget: Budget, sub_pdf: Callable[[bytes], Parsed]) -> Parsed:
    """Pages / Numbers / Keynote: the document format is private (IWA), so the embedded preview is read:
    QuickLook/Preview.pdf when present (older files), else the preview image (first page / slide / sheet)."""
    typ = {"pages": "document", "numbers": "spreadsheet", "key": "slides"}.get(ext, "document")
    lower = {n.lower(): n for n in pkg.names()}
    pdf_name = lower.get("quicklook/preview.pdf")
    if pdf_name:
        inner = sub_pdf(pkg.read(pdf_name, limit=80 * 1024 * 1024) or b"")
        inner.type = typ
        inner.counts["preview"] = "pdf"
        return inner
    for cand in ("preview.jpg", "quicklook/thumbnail.jpg", "preview-web.jpg", "quicklook/thumbnail.png", "preview.png"):
        name = lower.get(cand)
        if name:
            data = pkg.read(name, limit=40 * 1024 * 1024)
            marker = add_image(budget, data, "文稿预览图（第一页）", page=True) if data else None
            if marker:
                return Parsed(typ, f"（iWork 文稿，只能读取内嵌的第一页预览图）\n{marker}", "", {"preview": "image"})
    raise ParseError("unsupported", "iWork package without a preview")
