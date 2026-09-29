"""Small invented files of every type file-read handles, generated in the test (nothing is checked in).

All content here is invented for tests.
"""

from __future__ import annotations

import io
import plistlib
import struct
import zipfile
import zlib


def png(w: int = 320, h: int = 200, text: str = "", color=(250, 250, 250)) -> bytes:
    from PIL import Image, ImageDraw
    im = Image.new("RGB", (w, h), color)
    d = ImageDraw.Draw(im)
    d.rectangle([10, 10, w - 10, h - 10], outline=(0, 0, 0))
    if text:
        d.text((20, 20), text, fill=(0, 0, 0))
    buf = io.BytesIO()
    im.save(buf, "PNG")
    return buf.getvalue()


def _zip(files: dict, stored_first: str = "") -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        if stored_first:
            zf.writestr(zipfile.ZipInfo(stored_first), files[stored_first], compress_type=zipfile.ZIP_STORED)
        for name, data in files.items():
            if name != stored_first:
                zf.writestr(name, data)
    return buf.getvalue()


# ---- OOXML ---------------------------------------------------------------------------------------------

W_NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" ' \
       'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" ' \
       'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" ' \
       'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" ' \
       'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture"'


def _img_ext(data: bytes) -> str:
    return "png" if data.startswith(b"\x89PNG") else "jpg"


def _w_p(text: str, style: str = "") -> str:
    ppr = f'<w:pPr><w:pStyle w:val="{style}"/></w:pPr>' if style else ""
    return f"<w:p>{ppr}<w:r><w:t xml:space=\"preserve\">{text}</w:t></w:r></w:p>"


def docx(paragraphs: list, table: list | None = None, image: bytes | None = None, title: str = "") -> bytes:
    body = []
    for p in paragraphs:
        body.append(_w_p(*p) if isinstance(p, tuple) else _w_p(p))
    if table:
        rows = "".join("<w:tr>" + "".join(f"<w:tc>{_w_p(c)}</w:tc>" for c in row) + "</w:tr>" for row in table)
        body.append(f"<w:tbl>{rows}</w:tbl>")
    rels = ['<Relationship Id="rId9" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"'
            ' Target="http://example.invalid/x" TargetMode="External"/>']
    files = {}
    if image is not None:
        body.append('<w:p><w:r><w:drawing><wp:inline><a:graphic><a:graphicData><pic:pic><pic:blipFill>'
                    '<a:blip r:embed="rId5"/></pic:blipFill></pic:pic></a:graphicData></a:graphic></wp:inline>'
                    '</w:drawing></w:r></w:p>')
        rels.append('<Relationship Id="rId5" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                    f'relationships/image" Target="media/image1.{_img_ext(image)}"/>')
        files[f"word/media/image1.{_img_ext(image)}"] = image
    doc = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document {W_NS}><w:body>{"".join(body)}</w:body></w:document>'
    files["[Content_Types].xml"] = ('<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/'
                                    'content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-'
                                    'package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>'
                                    '<Default Extension="png" ContentType="image/png"/><Default Extension="jpg" '
                                    'ContentType="image/jpeg"/><Override PartName="/word/document.xml" ContentType="application/'
                                    'vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
    files["_rels/.rels"] = (f'<?xml version="1.0"?><Relationships {REL_NS}><Relationship Id="rId1" Type="http://schemas.'
                            'openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
                            '</Relationships>')
    files["word/document.xml"] = doc
    files["word/_rels/document.xml.rels"] = ('<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/'
                                             'package/2006/relationships">' + "".join(rels) + "</Relationships>")
    if title:
        files["docProps/core.xml"] = ('<?xml version="1.0"?><cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/'
                                      'package/2006/metadata/core-properties" xmlns:dc="http://purl.org/dc/elements/1.1/">'
                                      f"<dc:title>{title}</dc:title></cp:coreProperties>")
    files["word/vbaProject.bin"] = b"\x00MACRO-NEVER-READ"
    return _zip(files)


P_NS = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" ' \
       'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" ' \
       'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
REL_NS = 'xmlns="http://schemas.openxmlformats.org/package/2006/relationships"'


def chart_xml(title: str, series: dict, cats: list[str]) -> str:
    C = 'xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" ' \
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    sers = ""
    for name, vals in series.items():
        cat = "".join(f'<c:pt idx="{i}"><c:v>{c}</c:v></c:pt>' for i, c in enumerate(cats))
        val = "".join(f'<c:pt idx="{i}"><c:v>{v}</c:v></c:pt>' for i, v in enumerate(vals))
        sers += (f"<c:ser><c:tx><c:strRef><c:strCache><c:pt idx=\"0\"><c:v>{name}</c:v></c:pt></c:strCache></c:strRef></c:tx>"
                 f"<c:cat><c:strRef><c:strCache>{cat}</c:strCache></c:strRef></c:cat>"
                 f"<c:val><c:numRef><c:numCache>{val}</c:numCache></c:numRef></c:val></c:ser>")
    return (f'<c:chartSpace {C}><c:chart><c:title><c:tx><c:rich><a:p><a:r><a:t>{title}</a:t></a:r></a:p></c:rich></c:tx>'
            f'</c:title><c:plotArea><c:barChart>{sers}</c:barChart></c:plotArea></c:chart></c:chartSpace>')


def pptx(slides: list[dict], image: bytes | None = None) -> bytes:
    """slides: [{"title", "bullets": [...], "notes": str, "image": bool, "chart": chart_xml(...)}]"""
    files = {}
    ids, prels = [], []
    for i, s in enumerate(slides, 1):
        ids.append(f'<p:sldId id="{255 + i}" r:id="rId{i}"/>')
        prels.append(f'<Relationship Id="rId{i}" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                     f'relationships/slide" Target="slides/slide{i}.xml"/>')
        shapes = [f'<p:sp><p:nvSpPr><p:nvPr><p:ph type="title"/></p:nvPr></p:nvSpPr><p:txBody><a:p><a:r><a:t>'
                  f'{s["title"]}</a:t></a:r></a:p></p:txBody></p:sp>']
        paras = "".join(f"<a:p><a:r><a:t>{b}</a:t></a:r></a:p>" for b in s.get("bullets", []))
        shapes.append(f"<p:sp><p:nvSpPr><p:nvPr/></p:nvSpPr><p:txBody>{paras}</p:txBody></p:sp>")
        srels = []
        if s.get("image") and image is not None:
            shapes.append('<p:pic><p:blipFill><a:blip r:embed="rId2"/></p:blipFill></p:pic>')
            srels.append('<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                         f'relationships/image" Target="../media/image1.{_img_ext(image)}"/>')
            files[f"ppt/media/image1.{_img_ext(image)}"] = image
        if s.get("chart"):
            shapes.append('<p:graphicFrame><a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/'
                          '2006/chart"><c:chart xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
                          'r:id="rId4"/></a:graphicData></a:graphic></p:graphicFrame>')
            srels.append('<Relationship Id="rId4" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                         f'relationships/chart" Target="../charts/chart{i}.xml"/>')
            files[f"ppt/charts/chart{i}.xml"] = s["chart"]
        if s.get("notes"):
            srels.append(f'<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                         f'relationships/notesSlide" Target="../notesSlides/notesSlide{i}.xml"/>')
            files[f"ppt/notesSlides/notesSlide{i}.xml"] = (f'<p:notes {P_NS}><p:cSld><p:spTree><p:sp><p:txBody><a:p><a:r>'
                                                           f'<a:t>{s["notes"]}</a:t></a:r></a:p></p:txBody></p:sp>'
                                                           '</p:spTree></p:cSld></p:notes>')
        files[f"ppt/slides/slide{i}.xml"] = (f'<p:sld {P_NS}><p:cSld><p:spTree>{"".join(shapes)}</p:spTree>'
                                             '</p:cSld></p:sld>')
        files[f"ppt/slides/_rels/slide{i}.xml.rels"] = f'<Relationships {REL_NS}>{"".join(srels)}</Relationships>'
    files["ppt/presentation.xml"] = f'<p:presentation {P_NS}><p:sldIdLst>{"".join(ids)}</p:sldIdLst></p:presentation>'
    files["ppt/_rels/presentation.xml.rels"] = f'<Relationships {REL_NS}>{"".join(prels)}</Relationships>'
    overrides = "".join(f'<Override PartName="/ppt/slides/slide{i}.xml" ContentType="application/vnd.openxmlformats-'
                        'officedocument.presentationml.slide+xml"/>' for i in range(1, len(slides) + 1))
    files["[Content_Types].xml"] = ('<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/'
                                    'content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-'
                                    'package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>'
                                    '<Default Extension="png" ContentType="image/png"/><Default Extension="jpg" '
                                    'ContentType="image/jpeg"/><Override PartName="/ppt/presentation.xml" ContentType='
                                    '"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/>'
                                    f'{overrides}</Types>')
    files["_rels/.rels"] = (f'<?xml version="1.0"?><Relationships {REL_NS}><Relationship Id="rId1" Type="http://schemas.'
                            'openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="ppt/presentation.xml"/>'
                            '</Relationships>')
    return _zip(files)


def xlsx(sheets: dict) -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def huge_xlsx(rows: int, cols: int = 5) -> bytes:
    """A sheet written directly as XML (openpyxl would take minutes to write this many rows)."""
    def col(j: int) -> str:
        return chr(ord("A") + j)
    body = "".join(f'<row r="{r}">' + "".join(f'<c r="{col(j)}{r}"><v>{r * 10 + j}</v></c>' for j in range(cols))
                   + "</row>" for r in range(1, rows + 1))
    sheet = ('<?xml version="1.0" encoding="UTF-8"?><worksheet xmlns="http://schemas.openxmlformats.org/'
             f'spreadsheetml/2006/main"><dimension ref="A1:{col(cols - 1)}{rows}"/><sheetData>{body}</sheetData>'
             '</worksheet>')
    files = {
        "[Content_Types].xml": ('<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/'
                                'content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-'
                                'package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>'
                                '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-'
                                'officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/'
                                'sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.'
                                'worksheet+xml"/></Types>'),
        "_rels/.rels": (f'<?xml version="1.0"?><Relationships {REL_NS}><Relationship Id="rId1" Type="http://schemas.'
                        'openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
                        '</Relationships>'),
        "xl/workbook.xml": ('<?xml version="1.0"?><workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/'
                            'main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                            '<sheets><sheet name="Big" sheetId="1" r:id="rId1"/></sheets></workbook>'),
        "xl/_rels/workbook.xml.rels": (f'<?xml version="1.0"?><Relationships {REL_NS}><Relationship Id="rId1" Type='
                                       '"http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"'
                                       ' Target="worksheets/sheet1.xml"/></Relationships>'),
        "xl/worksheets/sheet1.xml": sheet,
    }
    return _zip(files)


# ---- OpenDocument / EPUB / iWork --------------------------------------------------------------------------

ODF_NS = ('xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
          'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
          'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
          'xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" '
          'xmlns:xlink="http://www.w3.org/1999/xlink"')


def odf(kind: str, body: str, manifest_extra: str = "") -> bytes:
    mime = {"odt": "application/vnd.oasis.opendocument.text", "ods": "application/vnd.oasis.opendocument.spreadsheet",
            "odp": "application/vnd.oasis.opendocument.presentation"}[kind]
    content = f'<?xml version="1.0"?><office:document-content {ODF_NS}><office:body>{body}</office:body></office:document-content>'
    manifest = ('<?xml version="1.0"?><manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0">'
                f'{manifest_extra}</manifest:manifest>')
    return _zip({"mimetype": mime, "content.xml": content, "META-INF/manifest.xml": manifest}, stored_first="mimetype")


def epub(title: str, author: str, chapters: list[str]) -> bytes:
    items = "".join(f'<item id="c{i}" href="c{i}.xhtml" media-type="application/xhtml+xml"/>' for i in range(len(chapters)))
    spine = "".join(f'<itemref idref="c{i}"/>' for i in range(len(chapters)))
    files = {
        "mimetype": "application/epub+zip",
        "META-INF/container.xml": ('<?xml version="1.0"?><container xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
                                   '<rootfiles><rootfile full-path="OEBPS/content.opf"/></rootfiles></container>'),
        "OEBPS/content.opf": ('<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" '
                              'xmlns:dc="http://purl.org/dc/elements/1.1/"><metadata>'
                              f"<dc:title>{title}</dc:title><dc:creator>{author}</dc:creator></metadata>"
                              f"<manifest>{items}</manifest><spine>{spine}</spine></package>"),
    }
    for i, c in enumerate(chapters):
        files[f"OEBPS/c{i}.xhtml"] = (f'<html xmlns="http://www.w3.org/1999/xhtml"><head><title>c{i}</title>'
                                      f"<style>p{{color:red}}</style></head><body><p>{c}</p></body></html>")
    return _zip(files, stored_first="mimetype")


def iwork_with_preview(image: bytes) -> bytes:
    return _zip({"Index/Document.iwa": b"\x00\x01private", "preview.jpg": image, "Metadata/Properties.plist": b""})


def iwork_with_pdf(pdf_bytes: bytes) -> bytes:
    return _zip({"index.xml": "<x/>", "QuickLook/Preview.pdf": pdf_bytes})


# ---- PDF ---------------------------------------------------------------------------------------------------


def text_pdf(pages: list[str]) -> bytes:
    """A minimal PDF with a Helvetica text layer (ASCII text)."""
    objs: list[bytes] = []
    kids = []
    n_pages = len(pages)
    # 1 catalog, 2 pages, 3 font, then (page, content) pairs
    for i, text in enumerate(pages):
        page_no = 4 + 2 * i
        kids.append(f"{page_no} 0 R")
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {n_pages} >>".encode())
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i, text in enumerate(pages):
        lines = text.split("\n")
        ops = ["BT /F1 14 Tf 72 720 Td 18 TL"]
        for ln in lines:
            esc = ln.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            ops.append(f"({esc}) Tj T*")
        ops.append("ET")
        stream = "\n".join(ops).encode("latin-1")
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> "
                    f"/Contents {5 + 2 * i} 0 R >>".encode())
        objs.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for n, body in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(f"{n} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def scanned_pdf(n_pages: int) -> bytes:
    from PIL import Image, ImageDraw
    ims = []
    for i in range(n_pages):
        im = Image.new("RGB", (850, 1100), (255, 255, 255))
        d = ImageDraw.Draw(im)
        d.text((80, 80), f"SCANNED PAGE {i + 1}", fill=(0, 0, 0))
        d.rectangle([60, 60, 790, 1040], outline=(0, 0, 0), width=3)
        ims.append(im)
    buf = io.BytesIO()
    ims[0].save(buf, "PDF", save_all=True, append_images=ims[1:], resolution=100)
    return buf.getvalue()


def encrypted_pdf() -> bytes:
    from pypdf import PdfReader, PdfWriter
    w = PdfWriter()
    for page in PdfReader(io.BytesIO(text_pdf(["Secret salary list"]))).pages:
        w.add_page(page)
    w.encrypt(user_password="pw-123", owner_password="owner-456")
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


# ---- OLE (as duck-typed stand-ins: olefile cannot write compound files) -------------------------------------


class FakeOle:
    def __init__(self, streams: dict[str, bytes]):
        self.streams = streams

    def exists(self, path: str) -> bool:
        return path in self.streams

    def openstream(self, path: str):
        return io.BytesIO(self.streams[path])

    def listdir(self, streams: bool = True, storages: bool = False):
        out = set()
        for p in self.streams:
            parts = p.split("/")
            if storages and len(parts) > 1:
                out.add((parts[0],))
            if streams:
                out.add(tuple(parts))
        return sorted(out)

    def close(self) -> None:
        pass


def doc_streams(text: str) -> dict:
    """WordDocument + 0Table with a one-piece table of 8-bit (cp1252) text."""
    body = text.encode("cp1252")
    wd = bytearray(0x400) + body
    struct.pack_into("<H", wd, 0, 0xA5EC)
    struct.pack_into("<H", wd, 0x0A, 0)          # 0Table, not encrypted
    fc = (0x400 * 2) | 0x40000000
    plc = struct.pack("<II", 0, len(body)) + struct.pack("<HIH", 0, fc, 0)
    clx = b"\x02" + struct.pack("<I", len(plc)) + plc
    table = b"\x00" * 16 + clx
    struct.pack_into("<II", wd, 0x01A2, 16, len(clx))
    return {"WordDocument": bytes(wd), "0Table": table}


def _rec(rtype: int, body: bytes, container: bool = False, inst: int = 0) -> bytes:
    ver_inst = (inst << 4) | (0xF if container else 0)
    return struct.pack("<HHI", ver_inst, rtype, len(body)) + body


def ppt_stream(slides: list[list[str]]) -> dict:
    inner = b""
    for texts in slides:
        inner += _rec(0x03F3, b"\x00" * 20)
        for t in texts:
            inner += _rec(0x0FA0, t.encode("utf-16-le"))
    slwt = _rec(0x0FF0, inner, container=True)
    doc = _rec(0x03E8, slwt, container=True)
    slides_recs = b"".join(_rec(0x03EE, _rec(0x03EF, b"\x00" * 8), container=True) for _ in slides)
    return {"PowerPoint Document": doc + slides_recs, "Current User": b""}


def msg_streams(subject: str, sender: str, to: str, body: str, attachment: tuple[str, bytes] | None = None) -> dict:
    enc = lambda s: s.encode("utf-16-le")  # noqa: E731
    streams = {"__substg1.0_0037001F": enc(subject), "__substg1.0_0C1A001F": enc(sender),
               "__substg1.0_0E04001F": enc(to), "__substg1.0_1000001F": enc(body),
               "__properties_version1.0": b"\x00" * 32}
    if attachment:
        streams["__attach_version1.0_#00000000/__substg1.0_3707001F"] = enc(attachment[0])
        streams["__attach_version1.0_#00000000/__substg1.0_37010102"] = attachment[1]
    return streams


def xls(rows: list[list]) -> bytes:
    import xlwt  # test-only dependency (BSD)
    wb = xlwt.Workbook()
    ws = wb.add_sheet("预算")
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            ws.write(r, c, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---- mail / calendar / contact / web -----------------------------------------------------------------------


def eml(subject: str, body: str, attachments: list[tuple[str, str, bytes]] = ()) -> bytes:
    from email.message import EmailMessage
    msg = EmailMessage()
    msg["Subject"] = subject
    msg["From"] = "周小满 <xiaoman@example.invalid>"
    msg["To"] = "李四 <lisi@example.invalid>"
    msg["Date"] = "Tue, 29 Sep 2026 09:30:00 +0800"
    msg.set_content(body)
    for name, mime, data in attachments:
        maintype, subtype = mime.split("/")
        msg.add_attachment(data, maintype=maintype, subtype=subtype, filename=name)
    return msg.as_bytes()


def ics() -> bytes:
    return ("BEGIN:VCALENDAR\r\nVERSION:2.0\r\nBEGIN:VEVENT\r\nUID:x1\r\nSUMMARY:季度复盘会\r\n"
            "DTSTART;TZID=Asia/Shanghai:20261015T143000\r\nDTEND;TZID=Asia/Shanghai:20261015T160000\r\n"
            "LOCATION:3楼大会议室\r\nORGANIZER;CN=王经理:mailto:wang@example.invalid\r\n"
            "ATTENDEE;CN=李四:mailto:lisi@example.invalid\r\nDESCRIPTION:请带上 Q3 数据\\n和改进计划\r\n"
            " （补充一行）\r\nEND:VEVENT\r\nEND:VCALENDAR\r\n").encode("utf-8")


def vcf() -> bytes:
    qp = "=E5=BC=A0=E4=B8=89"  # 张三
    return ("BEGIN:VCARD\r\nVERSION:2.1\r\nN;CHARSET=UTF-8;ENCODING=QUOTED-PRINTABLE:" + qp + ";;;\r\n"
            "FN;CHARSET=UTF-8;ENCODING=QUOTED-PRINTABLE:" + qp + "\r\nORG:青松物业\r\nTEL;CELL:139-0000-1234\r\n"
            "EMAIL:zhangsan@example.invalid\r\nEND:VCARD\r\n").encode("utf-8")


def webarchive(url: str, html: str) -> bytes:
    return plistlib.dumps({"WebMainResource": {"WebResourceURL": url, "WebResourceData": html.encode("utf-8"),
                                               "WebResourceMIMEType": "text/html",
                                               "WebResourceTextEncodingName": "UTF-8"}}, fmt=plistlib.FMT_BINARY)


def webloc(url: str) -> bytes:
    return plistlib.dumps({"URL": url}, fmt=plistlib.FMT_XML)


# ---- archives / bombs --------------------------------------------------------------------------------------


def zip_of(files: dict) -> bytes:
    return _zip(files)


def zip_bomb(uncompressed_mb: int = 1024) -> bytes:
    """One member declared (and actually) `uncompressed_mb` MiB of zeros, ~1 MB compressed."""
    buf = io.BytesIO()
    chunk = b"\x00" * (1024 * 1024)
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        with zf.open("zeros.txt", "w", force_zip64=True) as f:
            for _ in range(uncompressed_mb):
                f.write(chunk)
        zf.writestr("readme.txt", "hello from the bomb")
    return buf.getvalue()


def gzip_bomb(uncompressed_mb: int = 512) -> bytes:
    c = zlib.compressobj(9, zlib.DEFLATED, 31)
    chunk = b"\x00" * (1024 * 1024)
    out = [c.compress(chunk) for _ in range(uncompressed_mb)]
    out.append(c.flush())
    return b"".join(out)


def nested_zip(depth: int, leaf: bytes = b"deep secret text") -> bytes:
    data = _zip({"leaf.txt": leaf})
    for i in range(depth):
        data = _zip({f"level{i}.zip": data})
    return data


def mark_encrypted(zip_bytes: bytes) -> bytes:
    """Set the 'encrypted' flag on every member (stdlib cannot write encrypted zips)."""
    b = bytearray(zip_bytes)
    i = 0
    while True:
        i = b.find(b"PK\x03\x04", i)
        if i < 0:
            break
        b[i + 6] |= 0x1
        i += 4
    i = 0
    while True:
        i = b.find(b"PK\x01\x02", i)
        if i < 0:
            break
        b[i + 8] |= 0x1
        i += 4
    return bytes(b)
