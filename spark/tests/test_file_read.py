"""file-read: parsing every supported type in the sandbox, resource limits, the skill's validator, and the
kind=file item through ingest -> reading -> organizing -> /v1/state.

All data here is invented for tests; every file is generated in the test.
"""

from __future__ import annotations

import base64
import hashlib
import json
import subprocess
import sys
import time

import pytest

import filefixtures as F
from conftest import REPO, TEST_KEY, FakeChat, image_reader, ingest, make_item
from organizer import fileparse
from organizer.api import build_organizer
from organizer.clients import HashEmbedClient
from organizer.file_read import plain_summary, read_file
from organizer.fileparse import ole as O
from organizer.schemas import Item
from organizer.skills import SkillRegistry

REG = SkillRegistry(REPO / "skills")
VALIDATE = REG.script("file-read", "validate")


def parse(data: bytes, name: str) -> dict:
    return fileparse.parse_file(data, name)


# ---- routing + text per type ------------------------------------------------------------------------------


def test_docx_paragraphs_headings_tables_and_an_embedded_picture():
    data = F.docx([("物业通知", "Heading1"), "10 月 12 日（周日）8:00-18:00 全楼停水检修，请提前储水。"],
                  table=[["楼栋", "停水时段"], ["3 栋", "8:00-18:00"]], image=F.png(400, 300, "PUMP"),
                  title="停水通知")
    out = parse(data, "通知.docx")
    assert out["error"] is None and out["type"] == "document" and out["fmt"] == "docx"
    assert "# 物业通知" in out["text"] and "全楼停水检修" in out["text"]
    assert "| 3 栋 | 8:00-18:00 |" in out["text"]
    assert out["title"] == "停水通知"
    assert len(out["images"]) == 1 and "[[IMG:1]]" in out["text"]
    assert out["images"][0]["data"][:4] in (b"\x89PNG", b"\xff\xd8\xff\xe0")
    assert "MACRO" not in out["text"] and "example.invalid" not in out["text"]  # macros / external links unread


def test_pptx_slides_titles_bullets_notes_and_picture():
    data = F.pptx([{"title": "Q3 复盘", "bullets": ["营收 1,280 万", "新客 320 家"], "notes": "讲 5 分钟"},
                   {"title": "下一步", "bullets": ["10 月 20 日前定预算"], "image": True}], image=F.png(640, 360))
    out = parse(data, "复盘.pptx")
    assert out["type"] == "slides" and out["counts"]["slides"] == 2
    assert "## 第 1 张幻灯片：Q3 复盘" in out["text"] and "- 营收 1,280 万" in out["text"]
    assert "备注：讲 5 分钟" in out["text"] and "## 第 2 张幻灯片：下一步" in out["text"]
    assert len(out["images"]) == 1 and out["images"][0]["label"].startswith("第 2 张幻灯片")


def test_charts_in_slides_are_read_from_their_cached_data():
    chart = F.chart_xml("季度营收（万元）", {"2025": [120, 135.5], "2026": [150, 162]}, ["Q1", "Q2"])
    out = parse(F.pptx([{"title": "营收", "bullets": [], "chart": chart}]), "营收.pptx")
    assert "图表：季度营收（万元）\n| 类别 | 2025 | 2026 |" in out["text"]
    assert "| Q2 | 135.5 | 162 |" in out["text"]


def test_xlsx_sheets_become_compact_markdown_tables_with_names():
    import datetime as dt
    data = F.xlsx({"9月报销": [["姓名", "项目", "金额", "日期"], ["张三", "打车", 58.5, dt.date(2026, 9, 3)],
                              ["李四", "午餐", 120, dt.date(2026, 9, 4)]],
                   "汇总": [["合计", 178.5]]})
    out = parse(data, "报销.xlsx")
    assert out["type"] == "spreadsheet" and out["counts"]["sheets"] == 2
    assert "## 工作表：9月报销" in out["text"] and "## 工作表：汇总" in out["text"]
    assert "| 张三 | 打车 | 58.5 | 2026-09-03 |" in out["text"] and "| 李四 | 午餐 | 120 | 2026-09-04 |" in out["text"]


def test_spreadsheet_numbers_are_written_as_displayed():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "指标"
    ws.append(["转化率", "营收", "单价", "件数"])
    ws.append([0.2537, 1280, 36.5, 7])
    ws["A2"].number_format = "0.0%"
    ws["B2"].number_format = "#,##0.00"
    ws["C2"].number_format = '"¥"#,##0.00'
    buf = __import__("io").BytesIO()
    wb.save(buf)
    out = parse(buf.getvalue(), "指标.xlsx")
    assert "| 25.4% | 1,280.00 | ¥36.50 | 7 |" in out["text"]
    from organizer.fileparse.core import number_as_shown
    assert number_as_shown(1234.5, "[$$-409]#,##0.00") == "$1,234.50"
    assert number_as_shown(0.5, "0%") == "50%" and number_as_shown(3, "General") == "3"


def test_xls_legacy_workbook():
    out = parse(F.xls([["项目", "预算"], ["场地", 3000], ["餐饮", 1500.5]]), "预算.xls")
    assert out["error"] is None and out["type"] == "spreadsheet" and out["fmt"] == "xls"
    assert "## 工作表：预算" in out["text"] and "| 餐饮 | 1500.5 |" in out["text"]


def test_csv_in_gb18030_and_tsv():
    out = parse("姓名,电话\n王五,138\n".encode("gb18030"), "名单.csv")
    assert out["type"] == "spreadsheet" and "| 王五 | 138 |" in out["text"]
    out = parse(b"a\tb\n1\t2\n", "x.tsv")
    assert "| 1 | 2 |" in out["text"]


def test_text_pdf_and_scanned_pdf_pages_go_to_image_read():
    out = parse(F.text_pdf(["Invoice No. 04412233", "Total 1,280.00"]), "inv.pdf")
    assert out["type"] == "pdf" and out["counts"]["pages"] == 2 and not out["images"]
    assert "Invoice No. 04412233" in out["text"] and "## 第 2 页" in out["text"]
    out = parse(F.scanned_pdf(3), "scan.pdf")
    assert out["type"] == "scanned_pdf" and out["counts"]["pages"] == 3 and len(out["images"]) == 3
    assert [im["page"] for im in out["images"]] == [True] * 3 and "第 1 页" in out["images"][0]["label"]
    assert all(max(_size(im["data"])) <= 2560 for im in out["images"])


def test_scanned_pdf_reads_at_most_20_pages():
    out = parse(F.scanned_pdf(23), "long-scan.pdf")
    assert out["counts"]["pages"] == 23 and len(out["images"]) == 20
    assert any("上限" in n for n in out["notes"])


def _size(img: bytes) -> tuple[int, int]:
    import io
    from PIL import Image
    return Image.open(io.BytesIO(img)).size


def test_encrypted_pdf_is_reported_not_guessed():
    out = parse(F.encrypted_pdf(), "工资.pdf")
    assert out["error"] == "encrypted" and out["type"] == "pdf" and out["text"] == ""


def test_odt_ods_odp():
    out = parse(F.odf("odt", '<office:text><text:h text:outline-level="1">周报</text:h><text:p>本周完成<text:s/>3 项</text:p>'
                             '<text:list><text:list-item><text:p>联系供应商</text:p></text:list-item></text:list>'
                             '</office:text>'), "周报.odt")
    assert out["type"] == "document" and "# 周报" in out["text"] and "本周完成 3 项" in out["text"] and "- 联系供应商" in out["text"]
    # A sheet padded to a million rows / thousands of columns is never expanded.
    out = parse(F.odf("ods", '<office:spreadsheet><table:table table:name="库存"><table:table-row><table:table-cell>'
                             '<text:p>螺丝</text:p></table:table-cell><table:table-cell><text:p>200</text:p></table:table-cell>'
                             '<table:table-cell table:number-columns-repeated="16000"/></table:table-row>'
                             '<table:table-row table:number-rows-repeated="1048570"><table:table-cell '
                             'table:number-columns-repeated="16384"/></table:table-row></table:table></office:spreadsheet>'),
                "库存.ods")
    assert out["type"] == "spreadsheet" and "## 工作表：库存" in out["text"] and "| 螺丝 | 200 |" in out["text"]
    out = parse(F.odf("odp", '<office:presentation><draw:page><draw:frame><draw:text-box><text:p>开场</text:p>'
                             '</draw:text-box></draw:frame></draw:page><draw:page><text:p>结束</text:p></draw:page>'
                             '</office:presentation>'), "讲稿.odp")
    assert out["type"] == "slides" and out["counts"]["slides"] == 2 and "## 第 2 张幻灯片\n结束" in out["text"]


def test_encrypted_odf():
    out = parse(F.odf("odt", "<office:text/>", '<manifest:file-entry><manifest:encryption-data/></manifest:file-entry>'),
                "x.odt")
    assert out["error"] == "encrypted"


def test_epub_metadata_and_chapters():
    out = parse(F.epub("小王子", "圣埃克苏佩里", ["第一章 我六岁那年", "第二章 飞机坏了"]), "book.epub")
    assert out["type"] == "ebook" and "书名：小王子" in out["text"] and "第二章 飞机坏了" in out["text"]
    assert "color:red" not in out["text"]
    assert {"key": "title", "label": "书名", "value": "小王子"} in out["fields"]


def test_iwork_packages_read_their_preview():
    out = parse(F.iwork_with_preview(_jpeg()), "方案.pages")
    assert out["type"] == "document" and len(out["images"]) == 1 and "[[IMG:1]]" in out["text"]
    out = parse(F.iwork_with_pdf(F.text_pdf(["Budget 2027"])), "预算.numbers")
    assert out["type"] == "spreadsheet" and "Budget 2027" in out["text"]


def _jpeg() -> bytes:
    import io
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (800, 600), (200, 220, 240)).save(buf, "JPEG")
    return buf.getvalue()


def test_eml_headers_body_and_attachments_by_type():
    data = F.eml("报价确认", "李四你好，附件是 10 月报价，请周五前确认。",
                 [("报价.csv", "text/csv", "品名,单价\n打印纸,25\n".encode()), ("photo.png", "image/png", F.png(500, 400))])
    out = parse(data, "报价.eml")
    assert out["type"] == "email" and out["counts"]["attachments"] == 2
    assert out["text"].startswith("主题：报价确认\n发件人：周小满 <xiaoman@example.invalid>")
    assert "## 附件：报价.csv" in out["text"] and "| 打印纸 | 25 |" in out["text"]
    assert {a["filename"]: a["type"] for a in out["attachments"]} == {"报价.csv": "spreadsheet", "photo.png": "image"}
    assert {"key": "subject", "label": "主题", "value": "报价确认"} in out["fields"]
    assert len(out["images"]) == 1


def test_mbox_reads_the_first_20_messages():
    msgs = b"".join(b"From a@example.invalid Mon Sep 28 10:00:00 2026\n" + F.eml(f"第{i}封", f"正文{i}").replace(b"\r\n", b"\n")
                    + b"\n" for i in range(25))
    out = parse(msgs, "inbox.mbox")
    assert out["type"] == "email" and out["counts"]["messages"] == 25
    assert "主题：第19封" in out["text"] and "第20封" not in out["text"]


def test_msg_doc_ppt_from_ole_streams():
    # olefile cannot write compound files: the record walkers run on the same streams through a stand-in.
    fake = F.FakeOle(F.msg_streams("会议改期", "王经理", "李四", "会议改到周四 15:00。", ("议程.txt", "1. 预算\n".encode())))
    assert O.ole_kind(fake) == "msg"
    from organizer.fileparse.core import Budget
    from organizer.fileparse.dispatch import parse_bytes
    budget = Budget()
    p = O.parse_msg(fake, budget, lambda b, n: parse_bytes(b, n, "", budget, 1))
    assert p.type == "email" and "主题：会议改期" in p.text and "会议改到周四 15:00" in p.text
    assert p.attachments == [{"filename": "议程.txt", "type": "text", "summary": "1. 预算"}]
    doc = F.FakeOle(F.doc_streams("Lease renewal\rRent 4,500 per month\r"))
    assert O.ole_kind(doc) == "doc"
    p = O.parse_doc(doc)
    assert p.type == "document" and p.text == "Lease renewal\nRent 4,500 per month"
    ppt = F.FakeOle(F.ppt_stream([["年度计划", "目标 3 项"], ["时间表"]]))
    assert O.ole_kind(ppt) == "ppt"
    p = O.parse_ppt(ppt)
    assert p.type == "slides" and p.counts["slides"] == 2 and "## 第 2 张幻灯片\n时间表" in p.text
    enc = F.FakeOle({"EncryptionInfo": b"x", "EncryptedPackage": b"y"})
    assert O.ole_kind(enc) == "ooxml_encrypted"


def test_ics_and_vcf():
    out = parse(F.ics(), "invite.ics")
    assert out["type"] == "calendar"
    assert "日程：季度复盘会" in out["text"] and "开始：2026-10-15 14:30 (Asia/Shanghai)" in out["text"]
    assert "参与者：李四" in out["text"] and "（补充一行）" in out["text"]
    assert {"key": "location", "label": "地点", "value": "3楼大会议室"} in out["fields"]
    out = parse(F.vcf(), "张三.vcf")
    assert out["type"] == "contact" and "姓名：张三" in out["text"] and "电话：139-0000-1234" in out["text"]
    assert {"key": "phone", "label": "电话", "value": "139-0000-1234"} in out["fields"]


def test_web_pages_links_and_no_fetching():
    html = ("<html><head><title>装修攻略</title><script>fetch('http://evil.invalid')</script></head><body>"
            "<h1>厨房</h1><p>台面选石英石</p><img src='http://tracker.invalid/p.png'></body></html>")
    out = parse(html.encode(), "a.html")
    assert out["type"] == "web" and out["title"] == "装修攻略" and "# 厨房" in out["text"] and "fetch" not in out["text"]
    out = parse(F.webarchive("https://example.invalid/reno", html), "page.webarchive")
    assert out["type"] == "web" and out["text"].startswith("网址：https://example.invalid/reno") and "台面选石英石" in out["text"]
    out = parse(F.webloc("https://example.invalid/x"), "link.webloc")
    assert out["type"] == "web" and out["text"] == "链接：https://example.invalid/x"
    out = parse(b"[InternetShortcut]\r\nURL=https://example.invalid/y\r\n", "link.url")
    assert out["text"] == "链接：https://example.invalid/y"


def test_text_markdown_code_json_yaml_ipynb_rtf():
    assert parse("# 标题\n正文".encode(), "a.md")["title"] == "标题"
    out = parse(b"def total(xs):\n    return sum(xs)\n", "calc.py")
    assert out["type"] == "code" and "def total" in out["text"]
    out = parse(json.dumps({"城市": "杭州", "人数": 12}, ensure_ascii=False).encode(), "d.json")
    assert out["type"] == "data" and '"城市": "杭州"' in out["text"]
    out = parse(b"a: &x [1,2]\nb: *x\n", "c.yaml")
    assert out["type"] == "data" and out["text"] == "a: &x [1,2]\nb: *x"   # never loaded, never expanded
    nb = {"cells": [{"cell_type": "markdown", "source": ["# 分析"]},
                    {"cell_type": "code", "source": "print(1+1)", "outputs": [{"text": ["2\n"]}]}]}
    out = parse(json.dumps(nb).encode(), "n.ipynb")
    assert out["type"] == "code" and "print(1+1)" in out["text"] and "输出：2" in out["text"]
    rtf = rb"{\rtf1\ansi\ansicpg936{\fonttbl{\f0 SimSun;}}\f0 \'d6\'d0\'ce\'c4 and \u20320?\u22909? world\par second}"
    out = parse(rtf, "a.rtf")
    assert out["type"] == "document" and out["text"] == "中文 and 你好 world\nsecond"


def test_xmind_mind_maps_new_and_old_format():
    content = [{"title": "画布 1", "rootTopic": {"title": "开店计划", "children": {"attached": [
        {"title": "选址", "notes": {"plain": {"content": "看三个商圈"}}},
        {"title": "装修", "labels": ["10月"], "children": {"attached": [{"title": "吧台"}]}}]}}}]
    data = F.zip_of({"content.json": json.dumps(content, ensure_ascii=False).encode(), "metadata.json": b"{}",
                     "manifest.json": b"{}"})
    out = parse(data, "开店.xmind")
    assert out["type"] == "document" and out["fmt"] == "xmind"
    assert "## 画布：画布 1\n- 开店计划\n  - 选址\n    备注：看三个商圈\n  - 装修（10月）\n    - 吧台" in out["text"]
    xml8 = ('<xmap-content xmlns="urn:xmind:xmap:xmlns:content:2.0"><sheet><title>S</title><topic><title>年会</title>'
            '<children><topics type="attached"><topic><title>抽奖</title></topic></topics></children></topic></sheet>'
            '</xmap-content>')
    out = parse(F.zip_of({"content.xml": xml8, "meta.xml": "<meta/>"}), "old.xmind")
    assert "- 年会\n  - 抽奖" in out["text"]


def test_sqlite_databases_are_listed_from_memory():
    import sqlite3
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        path = f"{d}/x.db"
        c = sqlite3.connect(path)
        c.execute("CREATE TABLE 订单(id INTEGER, 客户 TEXT, 金额 REAL)")
        c.executemany("INSERT INTO 订单 VALUES (?,?,?)", [(i, f"客户{i}", i * 10.5) for i in range(30)])
        c.execute("CREATE VIEW v AS SELECT * FROM 订单")
        c.commit()
        c.close()
        data = open(path, "rb").read()
    out = parse(data, "shop.db")
    assert out["type"] == "data" and "## 表：订单（30 行）" in out["text"] and "| 1 | 客户1 | 10.5 |" in out["text"]
    assert "| 25 |" not in out["text"]  # first 20 rows only


def test_wps_office_files_route_like_office_97():
    from organizer.fileparse.core import sniff
    assert sniff(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64, "报告.wps") == "ole"
    assert parse(b"not an ole file at all", "表.et")["type"] == "spreadsheet"
    opml = b'<?xml version="1.0"?><opml><body><outline text="Q4 \xe8\xae\xa1\xe5\x88\x92"><outline text="A"/></outline></body></opml>'
    assert "outline：Q4 计划" in parse(opml, "plan.opml")["text"]


def test_xml_external_entities_are_refused():
    xxe = (b'<?xml version="1.0"?><!DOCTYPE r [<!ENTITY x SYSTEM "file:///etc/passwd">]><r>&x;</r>')
    out = parse(xxe, "evil.xml")
    assert out["error"] == "unsupported" and "root:" not in out["text"]
    lol = (b'<?xml version="1.0"?><!DOCTYPE l [<!ENTITY a "aaaaaaaaaa"><!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;&a;&a;">'
           b'<!ENTITY c "&b;&b;&b;&b;&b;&b;&b;&b;&b;&b;">]><l>&c;</l>')
    assert parse(lol, "lol.xml")["error"] == "unsupported"


def test_archives_read_each_entry_by_type():
    inner = F.zip_of({"notes.txt": "内层说明".encode()})
    data = F.zip_of({"报价.csv": "品名,单价\n胶带,3\n".encode(), "说明.md": "# 说明\n先看报价".encode(),
                     "__MACOSX/._x": b"junk", "inner.zip": inner, "song.mp3": b"ID3\x03fake"})
    out = parse(data, "资料.zip")
    assert out["type"] == "archive" and out["counts"]["entries"] == 4
    names = {a["filename"]: a for a in out["attachments"]}
    assert names["报价.csv"]["type"] == "spreadsheet" and names["inner.zip"]["type"] == "archive"
    # contract v6: audio / video entries are never read on the Spark; only their number is kept
    assert "song.mp3" not in names and out["media_skipped"] == 1
    assert "## 报价.csv" in out["text"] and "| 胶带 | 3 |" in out["text"] and "内层说明" in out["text"]
    out = parse(__import__("gzip").compress("压缩的日志".encode()), "app.log.gz")
    assert out["type"] == "text" and out["text"] == "压缩的日志"


def test_nesting_deeper_than_two_is_not_expanded():
    out = parse(F.nested_zip(4), "deep.zip")
    assert out["type"] == "archive" and "deep secret text" not in out["text"]


def test_encrypted_zip_and_unsupported_formats():
    assert parse(F.mark_encrypted(F.zip_of({"a.txt": b"secret"})), "locked.zip")["error"] == "encrypted"
    assert parse(b"7z\xbc\xaf\x27\x1c" + b"\x00" * 64, "a.7z")["error"] == "unsupported"
    assert parse(b"\x00\x00\x00\x18ftypheic" + b"\x00" * 64, "p.heic")["error"] == "unsupported"
    assert parse(b"\x00\x01\x02\x03\xff\xfe" * 100, "blob.bin")["error"] == "unsupported"
    assert parse(b"PK\x03\x04garbage", "broken.docx")["error"] == "corrupt"


# ---- resource limits ---------------------------------------------------------------------------------------


def test_zip_bomb_is_not_expanded():
    started = time.time()
    out = parse(F.zip_bomb(1024), "bomb.zip")
    assert time.time() - started < 30
    entries = {a["filename"]: a for a in out["attachments"]}
    assert entries["zeros.txt"]["summary"] == "过大，未读取" and "hello from the bomb" in out["text"]


def test_gzip_bomb_stops_at_the_archive_budget():
    out = parse(F.gzip_bomb(512), "zeros.gz")
    assert out["error"] == "too_large"


def test_huge_sheet_is_listed_partially_within_limits():
    started = time.time()
    out = parse(F.huge_xlsx(100_000, 5), "big.xlsx")
    assert out["error"] is None and time.time() - started < 60
    assert "## 工作表：Big" in out["text"] and "其余" in out["text"] and "未列出" in out["text"]
    assert out["text"].count("\n|") <= 402  # header + separator + 400 rows
    assert any("表格过大" in n for n in out["notes"])


def test_the_parser_process_is_killed_at_its_time_limit():
    out = fileparse.parse_file(F.scanned_pdf(3), "slow.pdf", wall=0.05)
    assert out["error"] == "too_large" and "秒" in out["notes"][0]


@pytest.mark.skipif(sys.platform != "linux", reason="RLIMIT_AS is enforced on Linux only")
def test_the_parser_process_has_a_memory_limit():
    # 80 MiB of address space is not enough to load PDFium and render a page; 30 MiB not enough to start.
    assert fileparse.parse_file(F.scanned_pdf(1), "p.pdf", mem=80 * 1024 * 1024)["error"] == "too_large"
    assert fileparse.parse_file(F.scanned_pdf(1), "p.pdf", mem=30 * 1024 * 1024)["error"] == "too_large"
    assert fileparse.parse_file(F.scanned_pdf(1), "p.pdf")["error"] is None


def test_the_parser_process_has_no_network():
    code = ("import sys; sys.path.insert(0, %r); from organizer.fileparse.worker import _no_network; _no_network()\n"
            "import socket\ntry:\n    socket.create_connection(('127.0.0.1', 9))\nexcept OSError as e:\n    print('refused', e)\n"
            "try:\n    socket.socket()\nexcept OSError as e:\n    print('refused', e)\n") % fileparse.SPARK_DIR
    out = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True, timeout=30).stdout
    assert out.count("refused") == 2


# ---- the skill's validator ---------------------------------------------------------------------------------

TEXT = "晨星文具有限公司\n发票号码：04412233\n开票日期：2026年09月18日\n价税合计：¥1,280.00\n打印纸 20 包"


def test_validator_accepts_grounded_summary_and_fields():
    out = {"summary": "晨星文具开出的发票，价税合计 1,280.00 元", "doc_kind": "receipt_invoice",
           "fields": [{"key": "merchant", "label": "", "value": "晨星文具有限公司"},
                      {"key": "doc_no", "label": "发票号码", "value": "04412233"},
                      {"key": "total", "label": "价税合计", "value": "¥1,280.00"}]}
    assert VALIDATE.validate(out, {"text": TEXT}) == []


def test_validator_rejects_invented_numbers_and_fields():
    out = {"summary": "这是一份发票，含税 1,132.74 元", "doc_kind": "receipt_invoice",
           "fields": [{"key": "tax", "label": "税额", "value": "147.26"},
                      {"key": "merchant", "label": "", "value": "晨星文具"},
                      {"key": "buyer", "label": "", "value": "未知"}]}
    errs = VALIDATE.validate(out, {"text": TEXT})
    assert any("1132.74" in e for e in errs) and any("开场" in e for e in errs)
    assert any("fields[0]" in e for e in errs) and any("fields[2]" in e and "占位" in e for e in errs)
    assert not any("fields[1]" in e for e in errs)  # a substring of the printed name is found in the text
    clean, dropped = VALIDATE.sanitize(out, {"text": TEXT})
    assert clean["summary"] == "" and [f["key"] for f in clean["fields"]] == ["merchant"] and dropped == 3


def test_validator_allows_two_parties_but_not_the_same_field_twice():
    text = "出租方（甲方）：许诺\n承租方（乙方）：吴昊"
    two = {"summary": "许诺出租给吴昊", "doc_kind": "contract",
           "fields": [{"key": "parties", "label": "出租方（甲方）", "value": "许诺"},
                      {"key": "parties", "label": "承租方（乙方）", "value": "吴昊"}]}
    assert VALIDATE.validate(two, {"text": text}) == []
    two["fields"][1]["value"] = "许诺"
    assert any("重复" in e for e in VALIDATE.validate(two, {"text": text}))


def test_validator_counts_and_non_receipt_kinds():
    assert VALIDATE.validate({"summary": "共 3 页的季度报告", "doc_kind": "report", "fields": []},
                             {"text": "季度报告", "extra_numbers": ["3"]}) == []
    errs = VALIDATE.validate({"summary": "季度报告", "doc_kind": "report",
                              "fields": [{"key": "title", "label": "", "value": "季度报告"}]}, {"text": "季度报告"})
    assert errs and "fields 写 []" in errs[0]


# ---- read_file: image parts, summary, fallbacks ---------------------------------------------------------------


def _harness(chat: FakeChat, tmp_path):
    from organizer.config import Settings
    s = Settings()
    s.data_dir = tmp_path / "d"
    s.skills_dir = REPO / "skills"
    s.start_worker = False
    s.embed_base_url = ""
    s.unlock_key = TEST_KEY
    return build_organizer(s, chat=chat, embedder=HashEmbedClient())


def test_read_file_inserts_image_readings_and_validates_the_summary(tmp_path):
    chat = FakeChat()
    chat.handlers["image-read"] = image_reader({"lines": ["水泵型号 WP-200", "扬程 32 米"], "gist": "水泵铭牌：WP-200"},
                                               "other")
    chat.push("file-read", {"summary": "物业通知 10 月 12 日停水，水泵型号 WP-200", "doc_kind": "notice", "fields": []})
    org = _harness(chat, tmp_path)
    data = F.docx(["10 月 12 日全楼停水检修。"], image=F.png(400, 300))
    # The Mac rebuilt this send copy with its pictures redacted (privacy review F3): only then are they read.
    res = read_file(org.harness, data, {"filename": "停水.docx", "mime": "", "source_app": {"name": "Finder"},
                                        "pictures_redacted": True}, subject="t")
    assert res.type == "document" and res.error is None and res.counts["images_read"] == 1
    assert "[文档图片 1·图片识别] 水泵铭牌：WP-200\n水泵型号 WP-200\n扬程 32 米" in res.text and "[[IMG" not in res.text
    assert res.summary == "物业通知 10 月 12 日停水，水泵型号 WP-200" and res.summary_source == "model"
    data_seen = [c[1] for c in chat.calls if c[0] == "file-read"][0]
    assert data_seen["filename"] == "停水.docx" and "水泵型号" in data_seen["text"]


def test_read_file_summary_with_a_made_up_number_falls_back_to_a_plain_summary(tmp_path):
    chat = FakeChat()
    bad = {"summary": "报价 999 元", "doc_kind": "other", "fields": []}
    chat.push("file-read", bad, bad)
    org = _harness(chat, tmp_path)
    res = read_file(org.harness, "报价单\n胶带 3 元".encode(), {"filename": "报价.txt"})
    assert res.summary_source == "plain" and res.summary == "文本：报价单"


def test_read_file_errors_use_local_text_and_a_plain_summary(tmp_path):
    chat = FakeChat()
    org = _harness(chat, tmp_path)
    res = read_file(org.harness, F.encrypted_pdf(), {"filename": "工资.pdf", "local_text": ""})
    assert res.error == "encrypted" and res.text == "" and res.summary == "PDF 文档「工资.pdf」已加密，需要密码，未读取内容"
    assert chat.count("file-read") == 0
    res = read_file(org.harness, None, {"filename": "旧合同.doc", "local_text": "租赁合同\n月租 4,500 元"})
    assert res.error is None and res.text == "租赁合同\n月租 4,500 元" and res.type == "document"
    wrong = read_file(org.harness, b"hello", {"filename": "a.txt", "sha256": "0" * 64})
    assert wrong.error == "corrupt"


def test_plain_summary_wording():
    assert plain_summary("spreadsheet", "a.xlsx", "", None) == "表格：a.xlsx"
    assert plain_summary("archive", "x.zip", "", "too_large") == "压缩包「x.zip」文件过大，未读取内容"


# ---- the item contract ------------------------------------------------------------------------------------------


def file_item(data: bytes | None, filename: str, *, minutes: int = 0, local_text: str | None = None, **kw) -> dict:
    item = make_item(None, kind="file", minutes=minutes, app="Finder", bundle="com.apple.finder", **kw)
    item["filename"] = filename
    item["mime"] = "application/octet-stream"
    if data is not None:
        item["bytes_b64"] = base64.b64encode(data).decode()
        item["size"] = len(data)
        item["sha256"] = hashlib.sha256(data).hexdigest()
    if local_text is not None:
        item["local_text"] = local_text
    item["captured_at"] = item.pop("started_at")
    item.pop("ended_at")
    return item


def test_item_contract_validation():
    ok = Item.model_validate(file_item(b"abc", "a.txt"))
    assert ok.kind == "file" and ok.started_at == ok.captured_at and ok.blob() == b"abc"
    with pytest.raises(ValueError):
        Item.model_validate({**file_item(b"abc", "a.txt"), "filename": ""})
    with pytest.raises(ValueError):
        Item.model_validate({**make_item("x"), "bytes_b64": base64.b64encode(b"a").decode()})
    with pytest.raises(ValueError):
        Item.model_validate({k: v for k, v in file_item(b"abc", "a.txt").items() if k != "bytes_b64"})
    assert Item.model_validate(file_item(None, "big.pdf", local_text="正文")).blob() is None
    big = file_item(b"", "huge.bin")
    big["bytes_b64"] = base64.b64encode(b"\x00" * (25 * 1024 * 1024 + 1)).decode()
    with pytest.raises(ValueError):
        Item.model_validate(big)


def test_file_item_is_read_organized_and_published_in_state(client, org, chat):
    chat.push("file-read", {"summary": "咖啡馆豆子报价，每公斤 120 元", "doc_kind": "receipt_invoice",
                            "fields": [{"key": "total", "label": "单价", "value": "120"}]})
    data = F.xlsx({"报价": [["品名", "单价"], ["咖啡馆豆子", 120]]})
    r = client.post("/v1/items", json={"items": [file_item(data, "咖啡馆报价.xlsx")]})
    assert r.status_code == 200 and r.json() == {"accepted": 1, "duplicates": 0}
    org.drain()
    state = client.get("/v1/state").json()
    [(item_id, reading)] = state["readings"].items()
    assert reading["source"] == "file-read" and reading["type"] == "spreadsheet"
    assert "| 咖啡馆豆子 | 120 |" in reading["text"] and reading["summary"] == "咖啡馆豆子报价，每公斤 120 元"
    assert reading["fields"] == [{"key": "total", "label": "单价", "value": "120"}]
    assert reading["counts"] == {"sheets": 1} and reading["attachments"] == [] and "error" not in reading
    assert reading["messages"] == [] and reading["numbers"] == []  # older clients' keys stay present
    [ev] = [e for e in state["events"] if item_id in e["item_ids"]]
    assert ev["title"] == "咖啡馆安排"  # the fake brief saw the reading text, not an empty item
    assign_item = [c[1]["item"] for c in chat.calls if c[0] == "event-assign"][0]
    assert assign_item["kind"] == "file" and assign_item["text"].startswith("咖啡馆豆子报价，每公斤 120 元\n文件：咖啡馆报价.xlsx")


def test_file_item_that_cannot_be_read_is_still_organized(client, org, chat):
    item = file_item(F.encrypted_pdf(), "工资.pdf", local_text="")
    assert client.post("/v1/items", json={"items": [item]}).status_code == 200
    org.drain()
    reading = client.get("/v1/state").json()["readings"][item["item_id"]]
    assert reading["error"] == "encrypted" and reading["summary"].startswith("PDF 文档「工资.pdf」已加密")
    assert chat.count("file-read") == 0 and chat.count("event-assign") == 1


def test_older_clients_items_and_state_are_unchanged(client, org):
    ingest(org, make_item("咖啡馆下周三开业"))
    org.drain()
    state = client.get("/v1/state").json()
    assert state["readings"] == {} and len(state["events"]) == 1


def test_a_long_document_is_split_on_its_reading_text(org, chat):
    text = ("咖啡馆装修进度：本周完成吧台，下周三验收。" * 3 + "\n\n" + "读书会十月书单确定为《活着》，周六下午见。" * 3)
    ingest(org, file_item(text.encode(), "周记.txt"))
    org.drain()
    assert chat.count("item-split") == 1
    parent = [r for r in org.store.all("SELECT item_id FROM items WHERE kind='file'")][0]["item_id"]
    segs = org.store.segments_of(parent)
    assert len(segs) == 2
    kids = [org.store.get_item(s["child_id"]) for s in segs]
    assert {k["kind"] for k in kids} == {"document"} and "吧台" in kids[0]["text"] and "活着" in kids[1]["text"]


# ---- video keyframes follow their media item -----------------------------------------------------------------


def test_video_keyframes_are_read_and_filed_with_their_recording(org, chat):
    chat.handlers["image-read"] = image_reader({"title": "咖啡馆开业方案", "subtitle": "", "bullets": [], "kpis": [],
                                                "footer": "", "page": "", "gist": "咖啡馆开业方案封面"}, "slide")
    video = make_item("咖啡馆开业会议录像的转写：周三验收吧台。", kind="imported_media", minutes=0)
    frame = make_item(None, kind="image", minutes=1, image_b64=base64.b64encode(F.png(640, 360)).decode())
    frame["parent_item_id"] = video["item_id"]
    frame["frame_ms"] = 12_000
    ingest(org, frame)          # the frame can arrive before the recording's transcript
    org.drain()
    assert org.store.current_event_link(frame["item_id"]) is None
    ingest(org, video)
    org.drain()
    ev = org.store.current_event_link(video["item_id"])["event_id"]
    assert org.store.current_event_link(frame["item_id"])["event_id"] == ev
    assert chat.count("event-assign") == 1  # the frame is never assigned on its own
    assert org.store.get_item(frame["item_id"])["meta"] == {"parent_item_id": video["item_id"], "frame_ms": 12000}
    # the user moves the recording: its frames follow
    from organizer.decisions import apply_decision
    other = make_item("读书会周六见", minutes=5)
    ingest(org, other)
    org.drain()
    to = org.store.current_event_link(other["item_id"])["event_id"]
    ok, _ = apply_decision(org, {"kind": "move_item", "item_id": video["item_id"], "to_event_id": to})
    assert ok
    org.drain()
    assert org.store.current_event_link(frame["item_id"])["event_id"] == to
